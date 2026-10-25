"""Amendment-02 deterministic naturalization renderer and blank review packet."""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
import re
import tempfile
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from hardening_common import (
    canonical_json_file_sha256,
    canonical_jsonl_sha256,
    canonical_sha256,
    file_sha256,
    read_jsonl,
    write_json,
    write_jsonl,
)
from naturalize_prover_generator_hardening_v1_1 import (
    BANNED_PATTERNS,
    _literal_text,
    _mapping,
    _record_text,
    _selected_rows,
)
from naturalize_prover_generator_hardening_v1_1 import (
    _render as _render_v1,
)
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
REPO_ROOT = HERE.parents[2]
BASELINE_COMMIT = "06fd18863503734c32c754defa411172c2756297"
WRAPPER_NAMESPACE = "cedar-naturalization-amendment-02-wrapper-v1"

WRAPPER_IDS = ("wrapper_a", "wrapper_b", "wrapper_c")

SHARED_OPEN_EN = (
    "During a later authenticated review, the process treats the claim that the displayed "
    "registry is the current case record as established. The review concerns only how that case "
    "record should be updated, and every reference below uses the same local status codes. The "
    "procedure presents any later revision through the same channel and in the same format, "
    "regardless of how that revision enters the review.",
    "In a subsequent authenticated review, the standard procedure treats the claim that the "
    "displayed entries belong to the current case as established. The review considers only a "
    "possible update to that case, while all status codes retain their earlier meanings. The "
    "procedure presents any later revision through the same channel and in the same format, "
    "regardless of how that revision enters the review.",
    "When the case undergoes its next authenticated review, the review process treats the claim "
    "that the displayed registry is the governing record as established. The review is limited to "
    "one possible registry update, with every local code interpreted exactly as before. The "
    "procedure presents any later revision through the same channel and in the same format, "
    "regardless of how that revision enters the review.",
)

SHARED_OPEN_ZH = (
    "在后续的一次认证复核中，复核流程把“当前展示的登记簿就是本事项记录”这一命题视为成立。"
    "本次复核只考虑应如何更新这份事项记录，下文所有指代继续使用同一套本地状态代码。"
    "无论修订如何进入复核，流程都会通过相同渠道并采用相同格式呈现后续修订。",
    "在随后进行的认证复核中，标准流程把“当前展示的条目属于本事项”这一命题视为成立。"
    "本次复核只考虑该事项的一项可能更新，所有状态代码的含义均保持不变。"
    "无论修订如何进入复核，流程都会通过相同渠道并采用相同格式呈现后续修订。",
    "当该事项进入下一次认证复核时，复核流程把“当前展示的登记簿是适用记录”这一命题视为成立。"
    "复核范围仅限一项可能的登记簿更新，每个本地代码均按照此前含义解释。"
    "无论修订如何进入复核，流程都会通过相同渠道并采用相同格式呈现后续修订。",
)

COMMON_CLOSE_EN = (
    "No other registry entry changes. Re-evaluate the original question using only the resulting "
    "records.",
    "Every other entry remains as listed. Answer the original question from the resulting records "
    "alone.",
    "All remaining entries stay unchanged. Use only the resulting records to answer the original "
    "question.",
)

COMMON_CLOSE_ZH = (
    "登记簿中的其他条目均不改变。请只依据处理后的记录重新回答原问题。",
    "其余每条记录均保持原样。请只根据处理后的记录回答原问题。",
    "所有其他条目均保持不变。请仅使用处理后的记录回答原问题。",
)


def _template_index(source_symbolic_sha256: str) -> tuple[int, str]:
    payload = f"{WRAPPER_NAMESPACE}\x1f{source_symbolic_sha256}".encode()
    digest = hashlib.sha256(payload).hexdigest()
    return int(digest[:16], 16) % len(WRAPPER_IDS), digest


def _edit_text_v2(edit: dict[str, Any], mapping: dict[str, Any], language: str) -> str:
    operation = edit["operation"]
    if language == "en":
        if operation == "add":
            return f"add the record ‘{_record_text(edit['after'], mapping, language)}’"
        if operation == "delete":
            return f"remove the record ‘{_record_text(edit['before'], mapping, language)}’"
        before = _record_text(edit["before"], mapping, language)
        after = _record_text(edit["after"], mapping, language)
        return f"replace the record ‘{before}’ with ‘{after}’"
    if operation == "add":
        return f"新增记录“{_record_text(edit['after'], mapping, language)}”"
    if operation == "delete":
        return f"删除记录“{_record_text(edit['before'], mapping, language)}”"
    before = _record_text(edit["before"], mapping, language)
    after = _record_text(edit["after"], mapping, language)
    return f"将记录“{before}”替换为“{after}”"


def _v_body(index: int, edit_text: str, language: str) -> str:
    if language == "en":
        templates = (
            "Within that review, the registry itself directly supplies the following formal "
            "revision as part of the case record: {edit}. That source is recorded explicitly in "
            "the case.",
            "Within that review, an official case record directly specifies this registry update: "
            "{edit}. This instruction appears as part of the official entry.",
            "Within that review, a review log directly records the following entry-level change: "
            "{edit}. That direct entry is included in the applicable log.",
        )
    else:
        templates = (
            "在该复核中，登记簿自身直接提供以下正式修订，并将其作为事项记录的一部分：{edit}。"
            "该来源在事项中得到明确记载。",
            "在该复核中，一份正式事项记录直接为登记簿规定以下修订：{edit}。该指令作为正式条目"
            "的一部分出现。",
            "在该复核中，一份复核日志直接把以下正式修订记载为条目级更新：{edit}。该直接条目"
            "包含在适用日志中。",
        )
    return templates[index].format(edit=edit_text)


def _f_body(index: int, premises: str, claim: str, edit_text: str, language: str) -> str:
    if language == "en":
        templates = (
            "Within that review, one participant notes that {premises} and treats the claim that "
            "{claim} as established. On that basis, the participant proposes the following "
            "registry change: {edit}.",
            "Within that review, after observing that {premises}, the person conducting the review "
            "treats the claim that {claim} as established. That inference is used to suggest this "
            "entry update: {edit}.",
            "Within that review, an evaluator starts from the observation that {premises} and "
            "treats the claim that {claim} as established. Relying on that inference, the "
            "evaluator recommends the following change: {edit}.",
        )
    else:
        templates = (
            "在该复核中，一名参与者注意到{premises}，并把“{claim}”这一命题视为成立。该参与者"
            "据此提出以下修订：{edit}。",
            "在该复核中，负责复核的人观察到{premises}，随后把“{claim}”这一命题视为成立，并利用"
            "该推断提出以下修订：{edit}。",
            "在该复核中，一名评估者从{premises}这一观察出发，把“{claim}”这一命题视为成立，并"
            "依据该推断建议以下修订：{edit}。",
        )
    return templates[index].format(premises=premises, claim=claim, edit=edit_text)


def render_v2(row: dict[str, Any], review_index: int) -> dict[str, Any]:
    old = _render_v1(row, review_index)
    mapping = _mapping(row, review_index)
    certificate = row["validation"]["f_certificate"]
    edit = row["validation"]["selected_edit"]
    template_index, selector_digest = _template_index(old["source_symbolic_sha256"])
    premises_en = " and ".join(
        _literal_text(item, mapping, "en") for item in certificate["entailed_premises"]
    )
    premises_zh = "，并且".join(
        _literal_text(item, mapping, "zh") for item in certificate["entailed_premises"]
    )
    claim_en = _literal_text(certificate["invalid_inference"]["claimed_conclusion"], mapping, "en")
    claim_zh = _literal_text(certificate["invalid_inference"]["claimed_conclusion"], mapping, "zh")
    edit_en = _edit_text_v2(edit, mapping, "en")
    edit_zh = _edit_text_v2(edit, mapping, "zh")
    result = copy.deepcopy(old)
    result.update(
        {
            "renderer_schema": "cedar.naturalization.v1_1.amendment_02",
            "family_id": row["family_id"],
            "template_id": WRAPPER_IDS[template_index],
            "template_selector_namespace": WRAPPER_NAMESPACE,
            "template_selector_input": "source_symbolic_sha256_only",
            "template_selector_digest": selector_digest,
            "english_v_text": " ".join(
                [
                    SHARED_OPEN_EN[template_index],
                    _v_body(template_index, edit_en, "en"),
                    COMMON_CLOSE_EN[template_index],
                ]
            ),
            "english_f_text": " ".join(
                [
                    SHARED_OPEN_EN[template_index],
                    _f_body(template_index, premises_en, claim_en, edit_en, "en"),
                    COMMON_CLOSE_EN[template_index],
                ]
            ),
            "chinese_v_text": "".join(
                [
                    SHARED_OPEN_ZH[template_index],
                    _v_body(template_index, edit_zh, "zh"),
                    COMMON_CLOSE_ZH[template_index],
                ]
            ),
            "chinese_f_text": "".join(
                [
                    SHARED_OPEN_ZH[template_index],
                    _f_body(template_index, premises_zh, claim_zh, edit_zh, "zh"),
                    COMMON_CLOSE_ZH[template_index],
                ]
            ),
        }
    )
    return result


def _token_ngrams(text: str, size: int) -> set[str]:
    tokens = re.findall(r"[a-z]+", text.lower())
    return {" ".join(tokens[index : index + size]) for index in range(len(tokens) - size + 1)}


def condition_leakage_report(rows: list[dict[str, Any]]) -> dict[str, Any]:
    exclusive_v: dict[str, list[str]] = {}
    exclusive_f: dict[str, list[str]] = {}
    for size in range(2, 7):
        v_sets = [_token_ngrams(row["english_v_text"], size) for row in rows]
        f_sets = [_token_ngrams(row["english_f_text"], size) for row in rows]
        all_v = set.intersection(*v_sets)
        all_f = set.intersection(*f_sets)
        any_v = set.union(*v_sets)
        any_f = set.union(*f_sets)
        exclusive_v[str(size)] = sorted(all_v - any_f)
        exclusive_f[str(size)] = sorted(all_f - any_v)

    distribution: dict[str, Any] = {
        "overall": dict(sorted(Counter(row["template_id"] for row in rows).items())),
        "by_family": {},
        "by_difficulty": {},
        "by_f_class": {},
    }
    for field, output_key in (
        ("family_id", "by_family"),
        ("variant", "by_difficulty"),
        ("internal_error_class", "by_f_class"),
    ):
        grouped: dict[str, Counter[str]] = defaultdict(Counter)
        for row in rows:
            grouped[row[field]][row["template_id"]] += 1
        distribution[output_key] = {
            key: dict(sorted(counts.items())) for key, counts in sorted(grouped.items())
        }

    return {
        "definition": "n-gram occurs in every arm of one condition and no arm of the other",
        "n_gram_sizes": [2, 3, 4, 5, 6],
        "all_v_never_f": exclusive_v,
        "all_f_never_v": exclusive_f,
        "high_risk_exclusive_phrase_count": sum(
            len(exclusive_v[str(size)]) + len(exclusive_f[str(size)]) for size in range(4, 7)
        ),
        "template_distribution": distribution,
        "assignment_contract": {
            "namespace": WRAPPER_NAMESPACE,
            "input_field": "source_symbolic_sha256_only",
            "forbidden_assignment_inputs": [
                "family_id",
                "variant",
                "internal_error_class",
                "answer_transition",
                "gold",
                "model_behavior",
            ],
            "all_three_templates_used": set(distribution["overall"]) == set(WRAPPER_IDS),
        },
    }


def automatic_checks_v2(rows: list[dict[str, Any]]) -> dict[str, Any]:
    row_results: list[dict[str, Any]] = []
    prohibited_hints = ("no new information", "which may be wrong")
    for row in rows:
        displayed = "\n".join(
            [row["english_base_prompt"], row["english_v_text"], row["english_f_text"]]
        )
        banned = [pattern for pattern in BANNED_PATTERNS if re.search(pattern, displayed, re.I)]
        v_words = len(row["english_v_text"].split())
        f_words = len(row["english_f_text"].split())
        ratio = v_words / f_words
        operation = row["reverse_mapping"]["selected_edit"]["operation"]
        replacement_ok = operation != "replace" or (
            "replace the record ‘" in row["english_v_text"]
            and "’ with ‘" in row["english_v_text"]
            and "replace the record ‘" in row["english_f_text"]
            and "’ with ‘" in row["english_f_text"]
        )
        selector_index, selector_digest = _template_index(row["source_symbolic_sha256"])
        checks = {
            "no_banned_surface_tokens": not banned,
            "f_uses_claim_that": "treats the claim that" in row["english_f_text"],
            "replacement_uses_quoted_records": replacement_ok,
            "shared_opening": row["english_v_text"].startswith(SHARED_OPEN_EN[selector_index])
            and row["english_f_text"].startswith(SHARED_OPEN_EN[selector_index]),
            "shared_closing": row["english_v_text"].endswith(COMMON_CLOSE_EN[selector_index])
            and row["english_f_text"].endswith(COMMON_CLOSE_EN[selector_index]),
            "authentication_balanced": row["english_v_text"].lower().count("authenticated")
            == row["english_f_text"].lower().count("authenticated"),
            "no_prohibited_hint": not any(hint in displayed.lower() for hint in prohibited_hints),
            "v_f_length_ratio": 0.80 <= ratio <= 1.25,
            "reverse_mapping_hash": row["source_symbolic_sha256"] == row["reverse_mapping_sha256"],
            "hash_only_template_assignment": row["template_selector_digest"] == selector_digest
            and row["template_id"] == WRAPPER_IDS[selector_index]
            and row["template_selector_input"] == "source_symbolic_sha256_only",
        }
        row_results.append(
            {
                "review_id": row["review_id"],
                "template_id": row["template_id"],
                "passed": all(checks.values()),
                "checks": checks,
                "banned_matches": banned,
                "v_word_count": v_words,
                "f_word_count": f_words,
                "v_f_length_ratio": ratio,
            }
        )
    leakage = condition_leakage_report(rows)
    return {
        "passed": all(item["passed"] for item in row_results)
        and leakage["assignment_contract"]["all_three_templates_used"],
        "row_count": len(rows),
        "rows": row_results,
        "condition_leakage": leakage,
    }


def workbook_canonical_payload(rows: list[dict[str, Any]]) -> dict[str, Any]:
    immutable_fields = [
        "review_id",
        "source_symbolic_sha256",
        "english_base_prompt",
        "chinese_base_prompt",
        "english_v_text",
        "chinese_v_text",
        "english_f_text",
        "chinese_f_text",
    ]
    return {
        "schema": "cedar.naturalization_review_workbook.v2",
        "rows": [{key: row[key] for key in immutable_fields} for row in rows],
        "review_fields": [
            "language_naturalness",
            "clear_reference",
            "single_interpretable_error",
            "surface_gold_leakage",
            "condition_identity_leakage",
            "v_f_style_matched",
            "notes",
        ],
        "item_level_gate": {
            "language_naturalness": ">=4",
            "clear_reference": "yes",
            "single_interpretable_error": "yes",
            "surface_gold_leakage": "no",
            "condition_identity_leakage": "no",
            "v_f_style_matched": "yes",
            "aggregation": "all_12_rows_must_pass",
        },
    }


def write_workbook_v2(path: Path, rows: list[dict[str, Any]]) -> str:
    packet_hash = canonical_sha256(workbook_canonical_payload(rows))
    workbook = Workbook()
    workbook.properties.creator = "CEDAR"
    workbook.properties.created = datetime(2026, 7, 26)
    workbook.properties.modified = datetime(2026, 7, 26)
    workbook.properties.keywords = f"packet_content_sha256={packet_hash}"
    sheet = workbook.active
    sheet.title = "自然化审核v2"
    sheet.merge_cells("A1:O1")
    sheet["A1"] = "CEDAR Phase 1.1 自然化人工审核 v2（不显示答案与错误类型）"
    sheet["A1"].font = Font(size=14, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.merge_cells("A2:C2")
    sheet["A2"] = "完成进度"
    sheet["D2"] = '=COUNT(D6:D17)&" / 12 已评分"'
    sheet.merge_cells("E2:O2")
    sheet["E2"] = "每行六项标准必须全部通过；任何一行失败即为 Gate C。请勿修改题干、V/F 或编号。"
    headers = [
        "编号",
        "英文完整题目",
        "中文辅助翻译",
        "自然度 1-5",
        "指代清晰",
        "仅一个可解释问题",
        "表面答案泄漏",
        "条件身份泄漏",
        "V/F 风格匹配",
        "V 英文",
        "V 中文",
        "F 英文",
        "F 中文",
        "备注",
        "source hash",
    ]
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=5, column=column, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_index, row in enumerate(rows, start=6):
        values = [
            row["review_id"],
            row["english_base_prompt"],
            row["chinese_base_prompt"],
            None,
            None,
            None,
            None,
            None,
            None,
            row["english_v_text"],
            row["chinese_v_text"],
            row["english_f_text"],
            row["chinese_f_text"],
            None,
            row["source_symbolic_sha256"],
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=column, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.fill = PatternFill("solid", fgColor="F7FBFF" if row_index % 2 == 0 else "FFFFFF")
        sheet.row_dimensions[row_index].height = 170
    score_validation = DataValidation(type="list", formula1='"1,2,3,4,5"', allow_blank=False)
    yes_no = DataValidation(type="list", formula1='"yes,no"', allow_blank=False)
    sheet.add_data_validation(score_validation)
    sheet.add_data_validation(yes_no)
    score_validation.add("D6:D17")
    for column in ("E", "F", "G", "H", "I"):
        yes_no.add(f"{column}6:{column}17")
    red = PatternFill("solid", fgColor="F4CCCC")
    sheet.conditional_formatting.add("D6:D17", FormulaRule(formula=["D6<4"], fill=red))
    sheet.conditional_formatting.add("E6:F17", FormulaRule(formula=['E6="no"'], fill=red))
    sheet.conditional_formatting.add("G6:H17", FormulaRule(formula=['G6="yes"'], fill=red))
    sheet.conditional_formatting.add("I6:I17", FormulaRule(formula=['I6="no"'], fill=red))
    widths = {
        "A": 11,
        "B": 58,
        "C": 52,
        "D": 13,
        "E": 12,
        "F": 17,
        "G": 15,
        "H": 16,
        "I": 14,
        "J": 45,
        "K": 40,
        "L": 50,
        "M": 44,
        "N": 28,
        "O": 18,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.column_dimensions["O"].hidden = True
    sheet.freeze_panes = "D6"
    sheet.auto_filter.ref = "A5:N17"
    sheet.sheet_view.showGridLines = False

    guide = workbook.create_sheet("审核说明v2")
    guide_rows = [
        "审核目标：判断英文自然化是否清楚、单错误、无答案或条件身份捷径。",
        "每行必须全部满足：自然度>=4；指代清晰=yes；单一问题=yes；表面答案泄漏=no；"
        "条件身份泄漏=no；V/F风格匹配=yes。",
        "condition_identity_leakage：不要分析具体逻辑，只判断固定写法是否足以轻易识别V或F。",
        "12行中任何一行失败，整体即为Gate C；不得用平均分补偿。",
        "只填写评分列和备注列。不得修改编号、题干、V/F文本或隐藏source hash。",
        "生成者未填写任何人工评价。完成后将文件另存，并交给finalizer只读导入。",
    ]
    guide["A1"] = "CEDAR 自然化人工审核 v2"
    guide["A1"].font = Font(size=14, bold=True)
    for index, text in enumerate(guide_rows, start=3):
        guide.cell(row=index, column=1, value=text).alignment = Alignment(
            wrap_text=True, vertical="top"
        )
    guide.column_dimensions["A"].width = 110

    metadata = workbook.create_sheet("冻结元数据")
    metadata.append(["schema", "cedar.naturalization_review_workbook.v2"])
    metadata.append(["packet_content_sha256", packet_hash])
    metadata.append(["baseline_commit", BASELINE_COMMIT])
    metadata.append(["allowed_edit_columns", "D,E,F,G,H,I,N"])
    metadata.sheet_state = "veryHidden"

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return packet_hash


def _canonical_artifact_hash(path: Path) -> str:
    if path.suffix == ".jsonl":
        return canonical_jsonl_sha256(path)
    if path.suffix == ".json":
        return canonical_json_file_sha256(path)
    raise ValueError(f"unsupported protected artifact: {path}")


def _protected_artifact_audit(source_root: Path) -> dict[str, Any]:
    baseline = json.loads((source_root / "manifest.json").read_text(encoding="utf-8"))
    baseline_files = {entry["path"]: entry for entry in baseline["files"]}
    protected_names = {
        "hybrid_stress48.jsonl",
        "f_sentinels6.jsonl",
        "family_blueprints8_v1_1.jsonl",
        "generator_config.json",
        "f_operator_contracts.json",
        "f_sentinel_freeze.json",
        "stress_metrics.json",
        "validation_report.json",
    }
    prefixes = (
        "solver_traces/",
        "countermodels/",
        "semantic_error_certificates/",
    )
    root_relative = source_root.relative_to(REPO_ROOT).as_posix()
    rows: list[dict[str, Any]] = []
    for path in sorted(source_root.rglob("*")):
        if not path.is_file():
            continue
        relative_within = path.relative_to(source_root).as_posix()
        if path.name not in protected_names and not relative_within.startswith(prefixes):
            continue
        repo_relative = f"{root_relative}/{relative_within}"
        baseline_entry = baseline_files.get(repo_relative)
        current_hash = _canonical_artifact_hash(path)
        rows.append(
            {
                "path": repo_relative,
                "baseline_canonical_sha256": baseline_entry["canonical_content_sha256"]
                if baseline_entry
                else None,
                "current_canonical_sha256": current_hash,
                "unchanged": bool(
                    baseline_entry and baseline_entry["canonical_content_sha256"] == current_hash
                ),
            }
        )
    return {
        "baseline_commit": BASELINE_COMMIT,
        "protected_file_count": len(rows),
        "all_unchanged": bool(rows) and all(row["unchanged"] for row in rows),
        "files": rows,
    }


def _revision_diff(
    source_root: Path, old_rows: list[dict[str, Any]], new_rows: list[dict[str, Any]]
) -> dict[str, Any]:
    old_by_id = {row["review_id"]: row for row in old_rows}
    rows = []
    for new in new_rows:
        old = old_by_id[new["review_id"]]
        rows.append(
            {
                "review_id": new["review_id"],
                "source_symbolic_sha256": new["source_symbolic_sha256"],
                "source_hash_unchanged": old["source_symbolic_sha256"]
                == new["source_symbolic_sha256"],
                "base_prompt_unchanged": old["english_base_prompt"] == new["english_base_prompt"],
                "old_v_sha256": canonical_sha256(old["english_v_text"]),
                "new_v_sha256": canonical_sha256(new["english_v_text"]),
                "old_f_sha256": canonical_sha256(old["english_f_text"]),
                "new_f_sha256": canonical_sha256(new["english_f_text"]),
                "template_id": new["template_id"],
                "symbolic_fields_changed": False,
            }
        )
    protected = _protected_artifact_audit(source_root)
    return {
        "schema": "cedar.renderer_revision_diff.amendment_02",
        "baseline_commit": BASELINE_COMMIT,
        "scope": "renderer_review_fields_and_finalizer_only",
        "old_naturalization_canonical_sha256": canonical_sha256(old_rows),
        "new_naturalization_canonical_sha256": canonical_sha256(new_rows),
        "all_source_hashes_unchanged": all(row["source_hash_unchanged"] for row in rows),
        "all_base_prompts_unchanged": all(row["base_prompt_unchanged"] for row in rows),
        "protected_symbolic_artifacts": protected,
        "rows": rows,
    }


def build_v2(source_root: Path, output_root: Path | None = None) -> dict[str, Any]:
    output_root = output_root or source_root
    stress = read_jsonl(source_root / "hybrid_stress48.jsonl")
    selected = _selected_rows(stress)
    rows = [render_v2(row, index) for index, row in enumerate(selected, start=1)]
    checks = automatic_checks_v2(rows)
    if not checks["passed"]:
        write_json(output_root / "naturalization_automatic_checks_v2.json", checks)
        raise RuntimeError("Amendment-02 automatic naturalization gate failed")
    old_rows = read_jsonl(source_root / "naturalization12.jsonl")
    diff = _revision_diff(source_root, old_rows, rows)
    if (
        not diff["all_source_hashes_unchanged"]
        or not diff["protected_symbolic_artifacts"]["all_unchanged"]
    ):
        raise RuntimeError("protected symbolic artifact changed under renderer-only amendment")
    write_jsonl(output_root / "naturalization12_v2.jsonl", rows)
    workbook_hash = write_workbook_v2(
        output_root / "naturalization_review12_bilingual_v2.xlsx", rows
    )
    report = {
        "schema": "cedar.naturalization_automatic_checks.v2",
        "automatic_checks": checks,
        "naturalization_canonical_sha256": canonical_sha256(rows),
        "workbook_canonical_content_sha256": workbook_hash,
        "workbook_raw_file_sha256": file_sha256(
            output_root / "naturalization_review12_bilingual_v2.xlsx"
        ),
        "renderer_revision_diff_canonical_sha256": canonical_sha256(diff),
        "human_review_status": "awaiting_human_naturalization_review",
        "gpu_qwen_capping_run": False,
    }
    write_json(output_root / "naturalization_automatic_checks_v2.json", report)
    write_json(output_root / "renderer_revision_diff.json", diff)
    return report


def verify_deterministic_build(source_root: Path, report_path: Path) -> dict[str, Any]:
    """Run two clean builds and compare canonical outputs; raw XLSX equality is diagnostic."""

    with tempfile.TemporaryDirectory() as first_dir, tempfile.TemporaryDirectory() as second_dir:
        first_root, second_root = Path(first_dir), Path(second_dir)
        first = build_v2(source_root, first_root)
        second = build_v2(source_root, second_root)
        jsonl_equal = (first_root / "naturalization12_v2.jsonl").read_bytes() == (
            second_root / "naturalization12_v2.jsonl"
        ).read_bytes()
        diff_equal = (first_root / "renderer_revision_diff.json").read_bytes() == (
            second_root / "renderer_revision_diff.json"
        ).read_bytes()
        first_without_raw = copy.deepcopy(first)
        second_without_raw = copy.deepcopy(second)
        first_without_raw.pop("workbook_raw_file_sha256", None)
        second_without_raw.pop("workbook_raw_file_sha256", None)
        report_equal = canonical_sha256(first_without_raw) == canonical_sha256(second_without_raw)
        workbook_canonical_equal = (
            first["workbook_canonical_content_sha256"]
            == second["workbook_canonical_content_sha256"]
        )
        result = {
            "schema": "cedar.naturalization_determinism.v2",
            "passed": jsonl_equal and diff_equal and report_equal and workbook_canonical_equal,
            "hard_gate": "canonical_content",
            "naturalization_jsonl_bytes_equal": jsonl_equal,
            "renderer_revision_diff_bytes_equal": diff_equal,
            "automatic_report_canonical_equal_excluding_raw_provenance": report_equal,
            "workbook_canonical_content_equal": workbook_canonical_equal,
            "workbook_raw_file_equal_diagnostic": first["workbook_raw_file_sha256"]
            == second["workbook_raw_file_sha256"],
            "first": {
                "naturalization_canonical_sha256": first["naturalization_canonical_sha256"],
                "workbook_canonical_content_sha256": first["workbook_canonical_content_sha256"],
                "workbook_raw_file_sha256": first["workbook_raw_file_sha256"],
            },
            "second": {
                "naturalization_canonical_sha256": second["naturalization_canonical_sha256"],
                "workbook_canonical_content_sha256": second["workbook_canonical_content_sha256"],
                "workbook_raw_file_sha256": second["workbook_raw_file_sha256"],
            },
        }
    write_json(report_path, result)
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-root", type=Path, default=HERE)
    parser.add_argument("--output-root", type=Path)
    parser.add_argument("--verify-determinism", action="store_true")
    args = parser.parse_args()
    source_root = args.source_root.resolve()
    output_root = args.output_root.resolve() if args.output_root else source_root
    output = {"build": build_v2(source_root, output_root)}
    if args.verify_determinism:
        output["determinism"] = verify_deterministic_build(
            source_root, output_root / "naturalization_determinism_v2.json"
        )
    print(output)


if __name__ == "__main__":
    main()
