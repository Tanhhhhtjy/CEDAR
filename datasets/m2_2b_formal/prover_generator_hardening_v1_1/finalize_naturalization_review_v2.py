"""Read-only import and deterministic Gate A/C finalization for a real v2 review workbook."""

from __future__ import annotations

import argparse
from pathlib import Path
from typing import Any

from hardening_common import canonical_sha256, file_sha256, read_jsonl, write_json
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
REPO_ROOT = HERE.parents[2]
DEFAULT_TEMPLATE = HERE / "naturalization_review12_bilingual_v2.xlsx"
DEFAULT_RESULT = HERE / "human_review_result.json"
DEFAULT_GATE = REPO_ROOT / "explore_outputs" / "prover_generator_hardening_v1_1" / "FINAL_GATE.json"

REVIEW_SHEET = "自然化审核v2"
ALLOWED_REVIEW_COLUMNS = {4, 5, 6, 7, 8, 9, 14}
REVIEW_ROWS = range(6, 18)


class ReviewPacketError(ValueError):
    """Raised before any output mutation when the completed workbook is invalid."""


def _cell_values(sheet: Any) -> list[list[Any]]:
    return [
        [sheet.cell(row=row, column=column).value for column in range(1, sheet.max_column + 1)]
        for row in range(1, sheet.max_row + 1)
    ]


def _verify_integrity(template_path: Path, completed_path: Path) -> dict[str, Any]:
    template = load_workbook(template_path, read_only=True, data_only=False)
    completed = load_workbook(completed_path, read_only=True, data_only=False)
    if completed.sheetnames != template.sheetnames:
        raise ReviewPacketError("sheet names or order changed")
    if completed.properties.keywords != template.properties.keywords:
        raise ReviewPacketError("frozen packet metadata changed")
    checked_cells = 0
    for sheet_name in template.sheetnames:
        expected = template[sheet_name]
        observed = completed[sheet_name]
        if (observed.max_row, observed.max_column) != (expected.max_row, expected.max_column):
            raise ReviewPacketError(f"sheet dimensions changed: {sheet_name}")
        for row in range(1, expected.max_row + 1):
            for column in range(1, expected.max_column + 1):
                if (
                    sheet_name == REVIEW_SHEET
                    and row in REVIEW_ROWS
                    and column in ALLOWED_REVIEW_COLUMNS
                ):
                    continue
                checked_cells += 1
                expected_value = expected.cell(row=row, column=column).value
                observed_value = observed.cell(row=row, column=column).value
                if observed_value != expected_value:
                    coordinate = observed.cell(row, column).coordinate
                    raise ReviewPacketError(f"immutable cell changed: {sheet_name}!{coordinate}")
    return {
        "passed": True,
        "immutable_cells_checked": checked_cells,
        "template_raw_sha256": file_sha256(template_path),
        "completed_raw_sha256": file_sha256(completed_path),
        "template_packet_metadata": template.properties.keywords,
    }


def _normalized_yes_no(value: Any, coordinate: str) -> str:
    if not isinstance(value, str) or value.strip().lower() not in {"yes", "no"}:
        raise ReviewPacketError(f"{coordinate} must be yes or no")
    return value.strip().lower()


def _normalized_score(value: Any, coordinate: str) -> int:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ReviewPacketError(f"{coordinate} must be an integer from 1 to 5")
    score = int(value)
    if score != value or not 1 <= score <= 5:
        raise ReviewPacketError(f"{coordinate} must be an integer from 1 to 5")
    return score


def _read_completed_rows(
    completed_path: Path, source_rows: list[dict[str, Any]]
) -> list[dict[str, Any]]:
    workbook = load_workbook(completed_path, read_only=True, data_only=False)
    sheet = workbook[REVIEW_SHEET]
    expected_by_id = {row["review_id"]: row for row in source_rows}
    results: list[dict[str, Any]] = []
    for row_number in REVIEW_ROWS:
        review_id = sheet.cell(row=row_number, column=1).value
        if review_id not in expected_by_id:
            raise ReviewPacketError(f"unknown or missing review ID at row {row_number}")
        expected = expected_by_id[review_id]
        source_hash = sheet.cell(row=row_number, column=15).value
        if source_hash != expected["source_symbolic_sha256"]:
            raise ReviewPacketError(f"source hash changed for {review_id}")
        score = _normalized_score(sheet.cell(row=row_number, column=4).value, f"D{row_number}")
        clear_reference = _normalized_yes_no(
            sheet.cell(row=row_number, column=5).value, f"E{row_number}"
        )
        single_error = _normalized_yes_no(
            sheet.cell(row=row_number, column=6).value, f"F{row_number}"
        )
        gold_leakage = _normalized_yes_no(
            sheet.cell(row=row_number, column=7).value, f"G{row_number}"
        )
        condition_leakage = _normalized_yes_no(
            sheet.cell(row=row_number, column=8).value, f"H{row_number}"
        )
        style_matched = _normalized_yes_no(
            sheet.cell(row=row_number, column=9).value, f"I{row_number}"
        )
        notes = sheet.cell(row=row_number, column=14).value
        if notes is not None and not isinstance(notes, str):
            raise ReviewPacketError(f"N{row_number} notes must be text or blank")
        checks = {
            "language_naturalness_gte_4": score >= 4,
            "clear_reference_yes": clear_reference == "yes",
            "single_interpretable_error_yes": single_error == "yes",
            "surface_gold_leakage_no": gold_leakage == "no",
            "condition_identity_leakage_no": condition_leakage == "no",
            "v_f_style_matched_yes": style_matched == "yes",
        }
        results.append(
            {
                "review_id": review_id,
                "source_symbolic_sha256": source_hash,
                "language_naturalness": score,
                "clear_reference": clear_reference,
                "single_interpretable_error": single_error,
                "surface_gold_leakage": gold_leakage,
                "condition_identity_leakage": condition_leakage,
                "v_f_style_matched": style_matched,
                "notes": notes,
                "checks": checks,
                "passed": all(checks.values()),
            }
        )
    if len(results) != 12 or len({row["review_id"] for row in results}) != 12:
        raise ReviewPacketError("all 12 unique review rows are required")
    return results


def gate_from_review_result(result: dict[str, Any]) -> dict[str, Any]:
    """Derive the frozen Gate A/C decision from a validated human-review result."""

    all_passed = result["status"] == "passed"
    common = {
        "symbolic_status": "symbolic_generator_passed",
        "human_review_status": "complete",
        "human_review_result_canonical_sha256": canonical_sha256(result),
        "gpu_qwen_capping_run": False,
    }
    if all_passed:
        return {
            "gate": "A",
            **common,
            "naturalization_status": "human_review_passed",
            "step_1_complete": True,
            "may_prepare_development_bank_protocol": True,
            "renderer_revision_only": False,
        }
    return {
        "gate": "C",
        **common,
        "naturalization_status": "naturalization_failed",
        "step_1_complete": False,
        "may_prepare_development_bank_protocol": False,
        "renderer_revision_only": True,
        "failed_review_ids": result["failed_review_ids"],
    }


def finalize_review(
    completed_path: Path,
    *,
    template_path: Path = DEFAULT_TEMPLATE,
    source_jsonl: Path | None = None,
    result_path: Path = DEFAULT_RESULT,
    final_gate_path: Path = DEFAULT_GATE,
) -> dict[str, Any]:
    """Validate without saving the workbook, then write the review result and Gate A/C."""

    completed_path = completed_path.resolve()
    template_path = template_path.resolve()
    source_jsonl = (source_jsonl or (template_path.parent / "naturalization12_v2.jsonl")).resolve()
    integrity = _verify_integrity(template_path, completed_path)
    source_rows = read_jsonl(source_jsonl)
    if len(source_rows) != 12:
        raise ReviewPacketError("frozen source JSONL must contain exactly 12 rows")
    rows = _read_completed_rows(completed_path, source_rows)
    all_passed = all(row["passed"] for row in rows)
    result = {
        "schema": "cedar.human_naturalization_review_result.v2",
        "status": "passed" if all_passed else "failed",
        "reviewed_row_count": 12,
        "complete": True,
        "integrity": integrity,
        "item_level_all_must_pass": True,
        "passed_row_count": sum(row["passed"] for row in rows),
        "failed_review_ids": [row["review_id"] for row in rows if not row["passed"]],
        "rows": rows,
        "result_canonical_sha256": canonical_sha256(rows),
        "agent_filled_human_ratings": False,
    }
    gate = gate_from_review_result(result)
    write_json(result_path.resolve(), result)
    write_json(final_gate_path.resolve(), gate)
    return {"result": result, "final_gate": gate}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("completed_workbook", type=Path)
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--source-jsonl", type=Path)
    parser.add_argument("--result", type=Path, default=DEFAULT_RESULT)
    parser.add_argument("--final-gate", type=Path, default=DEFAULT_GATE)
    args = parser.parse_args()
    try:
        output = finalize_review(
            args.completed_workbook,
            template_path=args.template,
            source_jsonl=args.source_jsonl,
            result_path=args.result,
            final_gate_path=args.final_gate,
        )
    except ReviewPacketError as error:
        print({"status": "invalid_review_packet", "error": str(error)})
        raise SystemExit(2) from error
    print(
        {
            "status": output["result"]["status"],
            "gate": output["final_gate"]["gate"],
            "failed_review_ids": output["result"]["failed_review_ids"],
        }
    )


if __name__ == "__main__":
    main()
