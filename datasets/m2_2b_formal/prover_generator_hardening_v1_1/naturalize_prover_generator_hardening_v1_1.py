"""Deterministic bilingual rendering and human-review workbook for 12 frozen records."""

from __future__ import annotations

import argparse
import copy
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from hardening_common import (
    ERROR_CLASSES,
    canonical_sha256,
    read_jsonl,
    write_json,
    write_jsonl,
)
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent

DOMAINS = [
    ("logistics registry", "shipment", "物流登记簿", "货件"),
    ("sample registry", "sample", "样本登记簿", "样本"),
    ("access registry", "account", "访问登记簿", "账户"),
    ("device registry", "device", "设备登记簿", "设备"),
    ("archive registry", "file", "档案登记簿", "档案"),
    ("routing registry", "node", "路由登记簿", "节点"),
    ("compliance registry", "case", "合规登记簿", "事项"),
    ("scheduling registry", "job", "调度登记簿", "任务"),
]

CODE_NAMES = [
    ("Amber", "琥珀"),
    ("Cobalt", "钴蓝"),
    ("Jade", "碧玉"),
    ("Ivory", "象牙"),
    ("Lilac", "丁香"),
    ("Copper", "赤铜"),
    ("Silver", "银灰"),
    ("Teal", "青绿"),
    ("Ochre", "赭黄"),
    ("Indigo", "靛蓝"),
    ("Coral", "珊瑚"),
    ("Pearl", "珍珠"),
    ("Umber", "棕褐"),
    ("Azure", "天蓝"),
    ("Saffron", "藏红"),
    ("Mint", "薄荷"),
    ("Slate", "石板"),
    ("Rose", "玫红"),
    ("Plum", "李紫"),
    ("Sand", "沙金"),
    ("Opal", "蛋白"),
    ("Navy", "海军蓝"),
    ("Lime", "青柠"),
    ("Ruby", "红玉"),
]

BANNED_PATTERNS = (
    r"\bf\d{2}\b",
    r"\bPGF\b",
    r"\?x",
    r"\btrue\b",
    r"\bfalse\b",
    r"\bunknown\b",
    r"\bvalid\b",
    r"\binvalid\b",
    r"\bcorrect\b",
    r"\bwrong\b",
    r"fatal error",
    r"No new information is provided",
    r"_",
)


def _selected_rows(stress: list[dict[str, Any]]) -> list[dict[str, Any]]:
    selected: list[dict[str, Any]] = []
    for error_class in ERROR_CLASSES:
        for variant in ("easy", "medium"):
            matches = sorted(
                (
                    row
                    for row in stress
                    if row["status"] == "valid"
                    and row["assigned_f_error_class"] == error_class
                    and row["variant"] == variant
                ),
                key=lambda row: row["attempt_id"],
            )
            if not matches:
                raise ValueError(f"no valid naturalization source for {error_class}/{variant}")
            selected.append(matches[0])
    return selected


def _mapping(row: dict[str, Any], review_index: int) -> dict[str, Any]:
    world = row["builder_output"]["initial_world"]
    predicates: dict[str, int] = {}
    for record in world["records"]:
        literals = (
            [record["literal"]] if record["kind"] == "fact" else record["body"] + [record["head"]]
        )
        for item in literals:
            predicates[item["predicate"]] = len(item["arguments"])
    query = row["builder_output"]["query"]
    predicates[query["predicate"]] = len(query["arguments"])
    certificate = row["validation"]["f_certificate"]
    for item in certificate["entailed_premises"] + [
        certificate["invalid_inference"]["claimed_conclusion"]
    ]:
        predicates[item["predicate"]] = len(item["arguments"])
    predicate_map = {}
    for index, predicate in enumerate(sorted(predicates)):
        en, zh = CODE_NAMES[index % len(CODE_NAMES)]
        predicate_map[predicate] = {
            "en": f"{en} {index + 1}",
            "zh": f"{zh}{index + 1}",
            "arity": predicates[predicate],
        }
    _, entity_noun, _, entity_noun_zh = DOMAINS[(review_index - 1) % len(DOMAINS)]
    constants = sorted(world["constants"])
    chinese_labels = ["甲", "乙", "丙", "丁", "戊", "己", "庚", "辛", "壬", "癸"]
    entity_map = {
        constant: {
            "en": f"{entity_noun.title()} {chr(65 + index)}",
            "zh": f"{entity_noun_zh}{chinese_labels[index]}",
        }
        for index, constant in enumerate(constants)
    }
    return {"predicates": predicate_map, "entities": entity_map}


def _argument_name(argument: str, mapping: dict[str, Any], language: str) -> str:
    if not argument.startswith("?"):
        return mapping["entities"][argument][language]
    variables_en = {"?x": "an item", "?y": "another item", "?z": "a third item"}
    variables_zh = {"?x": "某个对象", "?y": "另一个对象", "?z": "第三个对象"}
    return (variables_en if language == "en" else variables_zh).get(
        argument, "an item" if language == "en" else "某个对象"
    )


def _literal_text(item: dict[str, Any], mapping: dict[str, Any], language: str) -> str:
    code = mapping["predicates"][item["predicate"]][language]
    arguments = [_argument_name(arg, mapping, language) for arg in item["arguments"]]
    if language == "en":
        if len(arguments) == 1:
            core = f"{arguments[0]} has status {code}"
        else:
            core = f"{arguments[0]} has link {code} to {arguments[1]}"
        return f"it is not the case that {core}" if item["negated"] else core
    if len(arguments) == 1:
        core = f"{arguments[0]}具有{code}状态"
    else:
        core = f"{arguments[0]}与{arguments[1]}之间存在{code}连接"
    return f"并非{core}" if item["negated"] else core


def _record_text(record: dict[str, Any], mapping: dict[str, Any], language: str) -> str:
    if record["kind"] == "fact":
        return _literal_text(record["literal"], mapping, language)
    body = [_literal_text(item, mapping, language) for item in record["body"]]
    head = _literal_text(record["head"], mapping, language)
    if language == "en":
        return f"If {' and '.join(body)}, then {head}"
    return f"如果{'，并且'.join(body)}，那么{head}"


def _edit_text(edit: dict[str, Any], mapping: dict[str, Any], language: str) -> str:
    operation = edit["operation"]
    if language == "en":
        if operation == "add":
            return f"add the record that {_record_text(edit['after'], mapping, language)}"
        if operation == "delete":
            return f"remove the record that {_record_text(edit['before'], mapping, language)}"
        return (
            f"replace the record that {_record_text(edit['before'], mapping, language)} "
            f"with the record that {_record_text(edit['after'], mapping, language)}"
        )
    if operation == "add":
        return f"新增记录：{_record_text(edit['after'], mapping, language)}"
    if operation == "delete":
        return f"删除记录：{_record_text(edit['before'], mapping, language)}"
    return (
        f"将记录“{_record_text(edit['before'], mapping, language)}”替换为"
        f"“{_record_text(edit['after'], mapping, language)}”"
    )


def _render(row: dict[str, Any], review_index: int) -> dict[str, Any]:
    mapping = _mapping(row, review_index)
    domain_en, _, domain_zh, _ = DOMAINS[(review_index - 1) % len(DOMAINS)]
    world = row["builder_output"]["initial_world"]
    visible_records = world["records"]
    records_en = "\n".join(
        f"Record {index}. {_record_text(record, mapping, 'en')}."
        for index, record in enumerate(visible_records, start=1)
    )
    records_zh = "\n".join(
        f"记录{index}：{_record_text(record, mapping, 'zh')}。"
        for index, record in enumerate(visible_records, start=1)
    )
    query_en = _literal_text(row["builder_output"]["query"], mapping, "en")
    query_zh = _literal_text(row["builder_output"]["query"], mapping, "zh")
    question_en = (
        f"Is it supported that {query_en}? Choose one response: supported, contradicted, "
        "or undetermined."
    )
    question_zh = f"是否能够支持“{query_zh}”？请选择：支持、矛盾或无法确定。"
    base_en = (
        f"A closed {domain_en} uses local status codes. Treat every listed rule as exact and "
        "use only the records below. An unlisted statement is not automatically negated.\n\n"
        f"{records_en}\n\n{question_en}"
    )
    base_zh = (
        f"下面是一份封闭的{domain_zh}，其中使用本地状态代码。请严格按照所列规则，且只使用"
        f"下方记录；未列出的陈述不能自动视为否定。\n\n{records_zh}\n\n{question_zh}"
    )
    validation = row["validation"]
    edit = validation["selected_edit"]
    edit_en = _edit_text(edit, mapping, "en")
    edit_zh = _edit_text(edit, mapping, "zh")
    certificate = validation["f_certificate"]
    premises_en = " and ".join(
        _literal_text(item, mapping, "en") for item in certificate["entailed_premises"]
    )
    premises_zh = "，并且".join(
        _literal_text(item, mapping, "zh") for item in certificate["entailed_premises"]
    )
    claim_en = _literal_text(certificate["invalid_inference"]["claimed_conclusion"], mapping, "en")
    claim_zh = _literal_text(certificate["invalid_inference"]["claimed_conclusion"], mapping, "zh")
    common_end_en = "Now answer the original question using only the resulting records."
    common_end_zh = "现在请只依据处理后的记录回答原问题。"
    v_en = (
        f"During a later review, an authenticated case update instructs you to {edit_en}. "
        "This is the only authorized change to the registry; all other records remain as listed. "
        "The update supplies no further assumptions or changes. "
        f"{common_end_en}"
    )
    v_zh = (
        f"在后续复核中，一条经过认证的事项更新要求你{edit_zh}。这是登记簿中唯一获准的变更，"
        f"其他记录均保持原样。该更新不提供任何其他假设或变更。{common_end_zh}"
    )
    f_en = (
        f"During a later review, a reviewer notes that {premises_en}. The reviewer therefore "
        f"treats {claim_en} as established and, on that basis, instructs you to {edit_en}. "
        f"All other records remain as listed. {common_end_en}"
    )
    f_zh = (
        f"在后续复核中，一名复核者注意到{premises_zh}。该复核者因此把“{claim_zh}”视为已经"
        f"成立，并据此要求你{edit_zh}。其他记录均保持原样。{common_end_zh}"
    )
    source = {
        "initial_world": world,
        "query": row["builder_output"]["query"],
        "selected_edit": edit,
        "f_invalid_inference": certificate["invalid_inference"],
    }
    reverse = copy.deepcopy(source)
    return {
        "review_id": f"NAT-{review_index:02d}",
        "source_attempt_id": row["attempt_id"],
        "variant": row["variant"],
        "internal_error_class": row["assigned_f_error_class"],
        "english_base_prompt": base_en,
        "chinese_base_prompt": base_zh,
        "english_v_text": v_en,
        "chinese_v_text": v_zh,
        "english_f_text": f_en,
        "chinese_f_text": f_zh,
        "target_question_en": question_en,
        "target_question_zh": question_zh,
        "source_symbolic_sha256": canonical_sha256(source),
        "reverse_mapping": reverse,
        "reverse_mapping_sha256": canonical_sha256(reverse),
        "render_mapping_sha256": canonical_sha256(mapping),
    }


def automatic_checks(rows: list[dict[str, Any]]) -> dict[str, Any]:
    row_results: list[dict[str, Any]] = []
    for row in rows:
        displayed = "\n".join(
            [row["english_base_prompt"], row["english_v_text"], row["english_f_text"]]
        )
        banned = [pattern for pattern in BANNED_PATTERNS if re.search(pattern, displayed, re.I)]
        v_words = len(row["english_v_text"].split())
        f_words = len(row["english_f_text"].split())
        ratio = v_words / f_words
        checks = {
            "no_banned_surface_tokens": not banned,
            "sentence_terminators_present": all(
                text.strip().endswith((".", "?"))
                for text in (
                    row["english_base_prompt"],
                    row["english_v_text"],
                    row["english_f_text"],
                )
            ),
            "no_broken_nested_quotes": displayed.count('"') % 2 == 0,
            "v_f_target_identical": row["english_v_text"].endswith(
                "Now answer the original question using only the resulting records."
            )
            and row["english_f_text"].endswith(
                "Now answer the original question using only the resulting records."
            ),
            "v_f_length_ratio": 0.80 <= ratio <= 1.25,
            "reverse_mapping_hash": row["source_symbolic_sha256"] == row["reverse_mapping_sha256"],
        }
        row_results.append(
            {
                "review_id": row["review_id"],
                "passed": all(checks.values()),
                "checks": checks,
                "banned_matches": banned,
                "v_word_count": v_words,
                "f_word_count": f_words,
                "v_f_length_ratio": ratio,
            }
        )
    return {
        "passed": all(item["passed"] for item in row_results),
        "row_count": len(rows),
        "rows": row_results,
    }


def _workbook_canonical_payload(rows: list[dict[str, Any]]) -> dict[str, Any]:
    public_fields = [
        "review_id",
        "english_base_prompt",
        "chinese_base_prompt",
        "english_v_text",
        "chinese_v_text",
        "english_f_text",
        "chinese_f_text",
    ]
    return {
        "schema": "cedar.naturalization_review_workbook.v1",
        "rows": [{key: row[key] for key in public_fields} for row in rows],
        "review_fields": [
            "language_naturalness",
            "clear_reference",
            "single_interpretable_error",
            "surface_gold_leakage",
            "v_f_style_matched",
            "notes",
        ],
    }


def write_workbook(path: Path, rows: list[dict[str, Any]]) -> str:
    workbook = Workbook()
    workbook.properties.creator = "CEDAR"
    workbook.properties.created = datetime(2026, 7, 26)
    workbook.properties.modified = datetime(2026, 7, 26)
    sheet = workbook.active
    sheet.title = "自然化审核"
    sheet.merge_cells("A1:M1")
    sheet["A1"] = "CEDAR Phase 1.1 自然化人工审核（不显示答案与错误类型）"
    sheet["A1"].font = Font(size=14, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.merge_cells("A2:C2")
    sheet["A2"] = "完成进度"
    sheet["D2"] = '=COUNT(D6:D17)&" / 12 已评分"'
    sheet.merge_cells("E2:M2")
    sheet["E2"] = "请先读完整英文；中文仅作辅助。每行独立判断，不推测隐藏答案。"
    headers = [
        "编号",
        "英文完整题目",
        "中文辅助翻译",
        "自然度 1-5",
        "指代清晰",
        "仅一个可解释问题",
        "表面答案泄漏",
        "V/F 风格匹配",
        "V 英文",
        "V 中文",
        "F 英文",
        "F 中文",
        "备注",
    ]
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=5, column=column, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_index, row in enumerate(rows, start=6):
        values = [
            row["review_id"],
            row["english_base_prompt"],
            row["chinese_base_prompt"],
            "",
            "",
            "",
            "",
            "",
            row["english_v_text"],
            row["chinese_v_text"],
            row["english_f_text"],
            row["chinese_f_text"],
            "",
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=column, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.fill = PatternFill("solid", fgColor="F7FBFF" if row_index % 2 == 0 else "FFFFFF")
        sheet.row_dimensions[row_index].height = 165
    score_validation = DataValidation(type="list", formula1='"1,2,3,4,5"', allow_blank=True)
    yes_no = DataValidation(type="list", formula1='"yes,no"', allow_blank=True)
    sheet.add_data_validation(score_validation)
    sheet.add_data_validation(yes_no)
    score_validation.add("D6:D17")
    for column in ("E", "F", "G", "H"):
        yes_no.add(f"{column}6:{column}17")
    sheet.conditional_formatting.add(
        "D6:D17",
        FormulaRule(formula=["D6<3"], fill=PatternFill("solid", fgColor="F4CCCC")),
    )
    widths = {
        "A": 11,
        "B": 58,
        "C": 52,
        "D": 13,
        "E": 12,
        "F": 17,
        "G": 15,
        "H": 14,
        "I": 42,
        "J": 38,
        "K": 46,
        "L": 40,
        "M": 28,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "D6"
    sheet.auto_filter.ref = "A5:M17"
    sheet.sheet_view.showGridLines = False

    guide = workbook.create_sheet("审核说明")
    guide["A1"] = "审核目标"
    guide["A1"].font = Font(size=14, bold=True)
    guide["A3"] = "1. language_naturalness：英文是否自然、完整，1最低，5最高。"
    guide["A4"] = "2. clear_reference：实体、规则和后续复核中的指代是否清晰。"
    guide["A5"] = "3. single_interpretable_error：F段是否最多只有一个可解释的问题点。"
    guide["A6"] = "4. surface_gold_leakage：措辞是否直接暴露应选答案。"
    guide["A7"] = "5. v_f_style_matched：V与F在长度、语气和细节上是否大致匹配。"
    guide["A9"] = "请勿讨论或推测隐藏的形式标签；中文仅用于确认英文含义。"
    guide.column_dimensions["A"].width = 105
    for row in range(1, 10):
        guide[f"A{row}"].alignment = Alignment(wrap_text=True, vertical="top")
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return canonical_sha256(_workbook_canonical_payload(rows))


def build(root: Path) -> dict[str, Any]:
    stress = read_jsonl(root / "hybrid_stress48.jsonl")
    selected = _selected_rows(stress)
    rows = [_render(row, index) for index, row in enumerate(selected, start=1)]
    checks = automatic_checks(rows)
    if not checks["passed"]:
        write_json(root / "naturalization_automatic_checks.json", checks)
        raise RuntimeError("naturalization automatic gate failed")
    write_jsonl(root / "naturalization12.jsonl", rows)
    workbook_hash = write_workbook(root / "naturalization_review12_bilingual.xlsx", rows)
    result = {
        "automatic_checks": checks,
        "naturalization_canonical_sha256": canonical_sha256(rows),
        "workbook_canonical_content_sha256": workbook_hash,
        "human_review_status": "awaiting_human_naturalization_review",
    }
    write_json(root / "naturalization_automatic_checks.json", result)
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=HERE)
    args = parser.parse_args()
    print(build(args.root.resolve()))


if __name__ == "__main__":
    main()
