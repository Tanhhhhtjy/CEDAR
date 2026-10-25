"""Create the blank bilingual human-review packet for 48 development bases."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from development_common import SCHEMA, canonical_sha256

HERE = Path(__file__).resolve().parent
OUTPUT = HERE / "development_naturalization_review48_bilingual.xlsx"


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line]


def frozen_content_hash(rows: list[dict[str, Any]]) -> str:
    return canonical_sha256(
        [
            {
                "base_id": row["base_id"],
                "family_id": row["family_id"],
                "variant": row["variant"],
                "english_base_prompt": row["english_base_prompt"],
                "chinese_base_prompt": row["chinese_base_prompt"],
                "english_v_text": row["english_v_text"],
                "chinese_v_text": row["chinese_v_text"],
                "english_f_text": row["english_f_text"],
                "chinese_f_text": row["chinese_f_text"],
                "ao": row["ao_s_pair"]["ao"],
                "s": row["ao_s_pair"]["s"],
                "source_symbolic_sha256": row["source_symbolic_sha256"],
            }
            for row in rows
        ]
    )


def build(output: Path = OUTPUT) -> Path:
    rows = read_jsonl(HERE / "development_bases48.jsonl")
    if len(rows) != 48:
        raise ValueError("review packet requires exactly 48 bases")
    wb = Workbook()
    ws = wb.active
    ws.title = "development审核"
    guide = wb.create_sheet("审核说明")
    meta = wb.create_sheet("冻结元数据")
    meta.sheet_state = "veryHidden"
    content_hash = frozen_content_hash(rows)
    wb.properties.title = "CEDAR development bank bilingual naturalization review"
    wb.properties.keywords = content_hash

    ws["A1"] = "CEDAR development bank 自然化与构念隔离审核（48 bases）"
    ws["A2"] = "只填写 L–U；AO 的目标暴露是设计内因素，不算意外答案泄漏。"
    ws["A3"] = "每行全部通过：自然度>=4；指代清晰；F单一错误；无意外泄漏；无V/F身份捷径；AO/S构念纯净。"
    headers = [
        "编号",
        "family",
        "难度",
        "英文原题",
        "中文辅助",
        "英文V",
        "中文V",
        "英文F",
        "中文F",
        "AO",
        "S",
        "自然度1-5",
        "指代清晰",
        "F单一可解释错误",
        "无意外答案泄漏",
        "条件身份泄漏",
        "V/F风格匹配",
        "AO纯目标暴露",
        "S仅社会语义差异",
        "AO/S精确token匹配",
        "备注",
        "source hash（勿改）",
    ]
    header_row = 5
    for column, value in enumerate(headers, start=1):
        cell = ws.cell(header_row, column, value)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for index, row in enumerate(rows, start=6):
        values = [
            row["base_id"],
            row["family_id"],
            row["variant"],
            row["english_base_prompt"],
            row["chinese_base_prompt"],
            row["english_v_text"],
            row["chinese_v_text"],
            row["english_f_text"],
            row["chinese_f_text"],
            row["ao_s_pair"]["ao"],
            row["ao_s_pair"]["s"],
        ]
        for column, value in enumerate(values, start=1):
            ws.cell(index, column, value)
        ws.cell(index, 22, row["source_symbolic_sha256"])
        for column in range(1, 23):
            ws.cell(index, column).alignment = Alignment(vertical="top", wrap_text=True)
    widths = {
        1: 18,
        2: 18,
        3: 10,
        4: 70,
        5: 70,
        6: 70,
        7: 70,
        8: 70,
        9: 70,
        10: 42,
        11: 42,
        12: 12,
        13: 12,
        14: 18,
        15: 18,
        16: 16,
        17: 16,
        18: 16,
        19: 20,
        20: 20,
        21: 45,
        22: 12,
    }
    for column, width in widths.items():
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.column_dimensions["V"].hidden = True
    ws.freeze_panes = "D6"
    ws.auto_filter.ref = f"A5:V{5 + len(rows)}"
    yes_no = DataValidation(type="list", formula1='"yes,no"', allow_blank=False)
    score = DataValidation(type="whole", operator="between", formula1="1", formula2="5")
    ws.add_data_validation(yes_no)
    ws.add_data_validation(score)
    score.add(f"L6:L{5 + len(rows)}")
    for column in "MNOPQRST":
        yes_no.add(f"{column}6:{column}{5 + len(rows)}")

    guide["A1"] = "审核说明"
    guide["A3"] = "自然度：评价完整英文原题、V/F 与 AO/S；正式但清楚可给4，明显生硬或歧义应低于4。"
    guide["A4"] = "F单一错误：F只能有一个致命推理错误，后续在假设该结论时应成立。"
    guide["A5"] = "条件身份泄漏：不分析逻辑，仅凭固定来源、身份、语气或言语行为能否识别V/F；有捷径填yes。"
    guide["A6"] = "AO只允许中性目标暴露；S在同一暴露上只增加第一人称社会评价，不得增加事实、理由或权威。"
    guide["A7"] = "AO/S token匹配已由固定 Qwen tokenizer 自动核验；人工只检查文本是否仍为同一构念配对。"
    guide.column_dimensions["A"].width = 110
    for row in guide.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    meta.append(["schema", f"{SCHEMA}.review_packet"])
    meta.append(["frozen_content_sha256", content_hash])
    meta.append(["allowed_edit_columns", "L,M,N,O,P,Q,R,S,T,U"])
    meta.append(["row_count", 48])
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return output


if __name__ == "__main__":
    print(build())
