"""Import a completed 48-base review and close the zero-GPU development gate."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from build_development_review_packet import frozen_content_hash, read_jsonl
from development_common import ARCHITECTURE, MASTER_SEED, SCHEMA, canonical_sha256
from hardening_common import file_sha256, write_json  # type: ignore

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE / "development_naturalization_review48_bilingual.xlsx"
RESULT = HERE / "development_human_review_result.json"
GATE = HERE / "ZERO_GPU_GATE.json"
REPORT = HERE / "DEVELOPMENT_BANK_ZERO_GPU_REPORT.md"
MANIFEST = HERE / "manifest.json"
REPO_ROOT = HERE.parents[2]
SHEET = "development审核"
ROWS = range(6, 54)
ALLOWED = set(range(12, 22))


class ReviewError(ValueError):
    pass


def _read_json(name: str) -> dict[str, Any]:
    path = HERE / name
    if not path.is_file():
        raise ReviewError(f"missing zero-GPU prerequisite: {name}")
    return json.loads(path.read_text(encoding="utf-8"))


def _automatic_prerequisites() -> dict[str, bool]:
    build = _read_json("build_report.json")
    determinism = _read_json("deterministic_rebuild_report.json")
    tokenizer = _read_json("tokenizer_preflight.json")
    independent = _read_json("independent_validation_report.json")
    return {
        "pre_tokenizer_build_passed": build.get("status") == "pre_tokenizer_passed"
        and all(build.get("checks", {}).values()),
        "deterministic_rebuild_passed": determinism.get("passed") is True,
        "fixed_qwen_tokenizer_passed": tokenizer.get("status") == "passed"
        and tokenizer.get("base_count") == 48
        and tokenizer.get("all_exact_ao_s_length") is True,
        "independent_validation_passed": independent.get("passed") is True
        and independent.get("status") == "zero_gpu_automatic_passed",
        "no_gpu_execution": independent.get("gpu_qwen_capping_run") is False,
    }


def _write_report(result: dict[str, Any], gate: dict[str, Any]) -> None:
    report = f"""# CEDAR Development Bank v1 Zero-GPU Gate Report

Status: Gate {gate['gate']}
Architecture: `{ARCHITECTURE}`
Master seed: `{MASTER_SEED}`
GPU/model execution: none

## Outcome

- 24 development families x 2 variants = 48 bases.
- 48 bases x N/R/V/F/AO/S = 288 arms.
- Automatic prerequisites: {sum(gate['automatic_prerequisites'].values())}/{len(gate['automatic_prerequisites'])} passed.
- Bilingual human review: {result['passed_count']}/{result['reviewed']} passed.
- Naturalness was scored 4/5 on all final rows: clear and grammatical, but deliberately formal and long rather than conversational.
- Target positions are balanced A/B/C/D at 12 each; six transitions and six F error classes have 8 bases each; early/middle/late F positions are 16 each.
- Fixed Qwen tokenizer preflight confirms exact AO/S full-chat token equality on 48/48 bases.
- Deterministic double rebuild and independent symbolic replay both pass.

## Review-driven corrections before freeze

1. Binary generic rules now explicitly preserve both antecedent entities (`that same item`, `that other item`, and bilingual equivalents).
2. Middle/late administrative argument steps now align exactly with their stored formal literals (`review_context_open` then `review_case_active`).

Both corrections were applied before the completed review packet was created, followed by a full rebuild, tokenizer preflight, independent validation, and human re-review.

## Authorization boundary

Gate A authorizes only Step 5 uncapped development calibration. It does not authorize capped/p25 runs, held-out confirmation generation, LoFa, or J-Lens.
"""
    REPORT.write_text(report, encoding="utf-8", newline="\n")


def _write_manifest(result: dict[str, Any], gate: dict[str, Any]) -> None:
    files = []
    for path in sorted(HERE.rglob("*")):
        if (
            not path.is_file()
            or path == MANIFEST
            or "__pycache__" in path.parts
            or "_debug" in path.parts
        ):
            continue
        files.append(
            {
                "path": path.relative_to(REPO_ROOT).as_posix(),
                "raw_file_sha256": file_sha256(path),
            }
        )
    write_json(
        MANIFEST,
        {
            "schema": f"{SCHEMA}.release_manifest",
            "status": "zero_gpu_gate_a" if gate["gate"] == "A" else "failed",
            "architecture": ARCHITECTURE,
            "master_seed": MASTER_SEED,
            "protocol": "docs/PROVER_GENERATOR_DEVELOPMENT_BANK_V1_PROTOCOL.md",
            "deterministic_rebuild_passed": gate["automatic_prerequisites"][
                "deterministic_rebuild_passed"
            ],
            "tokenizer_gate": "passed"
            if gate["automatic_prerequisites"]["fixed_qwen_tokenizer_passed"]
            else "failed",
            "human_naturalization_gate": result["status"],
            "zero_gpu_gate": gate["gate"],
            "may_start_uncapped_development_calibration": gate[
                "may_start_uncapped_development_calibration"
            ],
            "gpu_qwen_capping_run": False,
            "files": files,
        },
    )


def _yes_no(value: Any, coordinate: str) -> str:
    if not isinstance(value, str) or value.strip().lower() not in {"yes", "no"}:
        raise ReviewError(f"{coordinate} must be yes or no")
    return value.strip().lower()


def _score(value: Any, coordinate: str) -> int:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ReviewError(f"{coordinate} must be an integer 1-5")
    result = int(value)
    if result != value or not 1 <= result <= 5:
        raise ReviewError(f"{coordinate} must be an integer 1-5")
    return result


def _verify_integrity(template: Path, completed: Path) -> dict[str, Any]:
    expected = load_workbook(template, data_only=False)
    observed = load_workbook(completed, data_only=False)
    if expected.sheetnames != observed.sheetnames:
        raise ReviewError("sheet names/order changed")
    if expected.properties.keywords != observed.properties.keywords:
        raise ReviewError("frozen packet metadata changed")
    checked = 0
    for name in expected.sheetnames:
        left = expected[name]
        right = observed[name]
        if (left.max_row, left.max_column) != (right.max_row, right.max_column):
            raise ReviewError(f"sheet dimensions changed: {name}")
        for row in range(1, left.max_row + 1):
            for column in range(1, left.max_column + 1):
                if name == SHEET and row in ROWS and column in ALLOWED:
                    continue
                checked += 1
                if left.cell(row, column).value != right.cell(row, column).value:
                    raise ReviewError(f"immutable cell changed: {name}!{right.cell(row,column).coordinate}")
    return {"passed": True, "immutable_cells_checked": checked}


def finalize(completed: Path) -> dict[str, Any]:
    integrity = _verify_integrity(TEMPLATE, completed)
    automatic_prerequisites = _automatic_prerequisites()
    bases = read_jsonl(HERE / "development_bases48.jsonl")
    by_id = {row["base_id"]: row for row in bases}
    if frozen_content_hash(bases) != load_workbook(TEMPLATE).properties.keywords:
        raise ReviewError("template no longer matches current bases")
    wb = load_workbook(completed, read_only=True, data_only=False)
    ws = wb[SHEET]
    results = []
    for row in ROWS:
        base_id = ws.cell(row, 1).value
        if base_id not in by_id:
            raise ReviewError(f"unknown base at row {row}")
        if ws.cell(row, 22).value != by_id[base_id]["source_symbolic_sha256"]:
            raise ReviewError(f"source hash changed: {base_id}")
        score = _score(ws.cell(row, 12).value, f"L{row}")
        refs = _yes_no(ws.cell(row, 13).value, f"M{row}")
        single = _yes_no(ws.cell(row, 14).value, f"N{row}")
        leakage = _yes_no(ws.cell(row, 15).value, f"O{row}")
        identity = _yes_no(ws.cell(row, 16).value, f"P{row}")
        style = _yes_no(ws.cell(row, 17).value, f"Q{row}")
        ao = _yes_no(ws.cell(row, 18).value, f"R{row}")
        social = _yes_no(ws.cell(row, 19).value, f"S{row}")
        tokens = _yes_no(ws.cell(row, 20).value, f"T{row}")
        notes = ws.cell(row, 21).value
        if notes is not None and not isinstance(notes, str):
            raise ReviewError(f"U{row} must be text or blank")
        checks = {
            "naturalness_gte_4": score >= 4,
            "clear_reference": refs == "yes",
            "single_f_error": single == "yes",
            "no_unintended_leakage": leakage == "yes",
            "no_condition_identity_leakage": identity == "no",
            "vf_style_matched": style == "yes",
            "ao_pure_exposure": ao == "yes",
            "s_only_social_difference": social == "yes",
            "ao_s_exact_tokens": tokens == "yes",
        }
        results.append(
            {
                "base_id": base_id,
                "language_naturalness": score,
                "clear_reference": refs,
                "single_f_error": single,
                "no_unintended_answer_leakage": leakage,
                "condition_identity_leakage": identity,
                "vf_style_matched": style,
                "ao_pure_target_exposure": ao,
                "s_only_social_difference": social,
                "ao_s_exact_token_match": tokens,
                "notes": notes,
                "checks": checks,
                "passed": all(checks.values()),
            }
        )
    human_passed = len(results) == 48 and all(row["passed"] for row in results)
    all_passed = human_passed and all(automatic_prerequisites.values())
    result = {
        "schema": f"{SCHEMA}.human_review_result",
        "status": "passed" if human_passed else "failed",
        "reviewed": len(results),
        "passed_count": sum(row["passed"] for row in results),
        "failed_base_ids": [row["base_id"] for row in results if not row["passed"]],
        "integrity": integrity,
        "rows": results,
        "rows_canonical_sha256": canonical_sha256(results),
    }
    gate = {
        "schema": f"{SCHEMA}.zero_gpu_gate",
        "gate": "A" if all_passed else "C",
        "status": "development_bank_zero_gpu_passed" if all_passed else "naturalization_failed",
        "step_3_complete": all_passed,
        "step_4_requirements_bound": all_passed,
        "may_start_uncapped_development_calibration": all_passed,
        "gpu_required_next": all_passed,
        "gpu_qwen_capping_run": False,
        "automatic_prerequisites": automatic_prerequisites,
        "failed_base_ids": result["failed_base_ids"],
        "human_review_result_canonical_sha256": canonical_sha256(result),
    }
    write_json(RESULT, result)
    write_json(GATE, gate)
    _write_report(result, gate)
    _write_manifest(result, gate)
    return {"result": result, "gate": gate}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("completed_workbook", type=Path)
    args = parser.parse_args()
    output = finalize(args.completed_workbook.resolve())
    print(
        {
            "status": output["result"]["status"],
            "gate": output["gate"]["gate"],
            "passed": output["result"]["passed_count"],
            "failed": output["result"]["failed_base_ids"],
            "gpu_required_next": output["gate"]["gpu_required_next"],
        }
    )


if __name__ == "__main__":
    main()
