#!/usr/bin/env python3
"""Independent CPU validator for the E2 new-salt bank and review boundary."""

from __future__ import annotations

import argparse
import json
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from build_human_review_workbook import (
    FIRST_REVIEW_ROW,
    LAST_REVIEW_ROW,
    SHEET_NAME,
    WORKBOOK_NAME,
    visible_payload,
)
from e2_new_salt_common import (
    ARMS,
    ERROR_CLASSES,
    HERE,
    MODEL_REPO,
    MODEL_REVISION,
    OPTION_LETTERS,
    ORDERED_TRANSITIONS,
    SALT_SHA256,
    SCHEMA,
    canonical_jsonl_sha256,
    canonical_sha256,
    canonical_text_sha256,
    file_sha256,
    object_sha256,
    read_json,
    require,
    write_json,
)
from e2_new_salt_run_common import selection_fingerprint
from openpyxl import load_workbook

PARENT_PROVER9_REFERENCE_SHA256 = (
    "53747e12206b94e3be7f9a464104d38b5461adf1ffb8d06ce4af2de63063f80c"
)
PARENT_MACE4_REFERENCE_SHA256 = (
    "868516e4c938fb6bb2b9135ba9b887c49bb6d9460b1e31a1b36cfa06814b31f7"
)


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line]


def protocol_hashes() -> dict[str, str]:
    freeze = read_json(HERE.parents[2] / "docs" / "E2_PROTOCOL_FREEZE_MANIFEST.json")
    observed = {
        relative: canonical_text_sha256(HERE.parents[2] / relative)
        for relative in freeze["files"]
    }
    require(observed == freeze["files"], "Frozen E2 protocol canonical hashes changed")
    return observed


def validate_workbook(bases: list[dict[str, Any]]) -> dict[str, Any]:
    packet = read_json(HERE / "HUMAN_REVIEW_PACKET.json")
    workbook_path = HERE / WORKBOOK_NAME
    workbook = load_workbook(workbook_path, read_only=False, data_only=False)
    require(workbook.sheetnames == [SHEET_NAME, "审核说明", "审核声明"], "Workbook sheets changed")
    require(
        workbook.properties.keywords == packet["workbook_canonical_content_sha256"],
        "Workbook packet metadata changed",
    )
    sheet = workbook[SHEET_NAME]
    blank_cells = []
    for row in range(FIRST_REVIEW_ROW, LAST_REVIEW_ROW + 1):
        for column in range(8, 17):
            if sheet.cell(row, column).value is not None:
                blank_cells.append(sheet.cell(row, column).coordinate)
    attest = workbook["审核声明"]
    for coordinate in ("B2", "B3", "B4", "B5"):
        if attest[coordinate].value is not None:
            blank_cells.append(f"审核声明!{coordinate}")
    immutable_text = "\n".join(
        str(sheet.cell(row, column).value or "")
        for row in range(1, sheet.max_row + 1)
        for column in range(1, 8)
    ).casefold()
    forbidden = (
        "error_class",
        "symbolic_sha",
        "prover9",
        "mace4",
        "qwen",
        "full_p25",
        "gold_answer",
    )
    canonical_payload = object_sha256(visible_payload(bases))
    checks = {
        "rows_24": LAST_REVIEW_ROW - FIRST_REVIEW_ROW + 1 == 24,
        "ratings_blank": not blank_cells,
        "canonical_content_hash": canonical_payload
        == packet["workbook_canonical_content_sha256"],
        "raw_file_hash": file_sha256(workbook_path) == packet["workbook_raw_sha256"],
        "no_forbidden_internal_markers": not any(token in immutable_text for token in forbidden),
        "gold_hidden": packet["blinding"]["gold_hidden"] is True,
        "f_class_hidden": packet["blinding"]["f_class_hidden"] is True,
        "symbolic_internals_hidden": packet["blinding"]["symbolic_internals_hidden"] is True,
        "qwen_outputs_absent": packet["blinding"]["qwen_outputs_absent"] is True,
        "agent_ratings_absent": packet["agent_filled_human_ratings"] is False,
    }
    return {
        "status": "passed" if all(checks.values()) else "failed",
        "checks": checks,
        "unexpected_filled_cells": blank_cells,
        "workbook_raw_sha256": packet["workbook_raw_sha256"],
        "workbook_canonical_content_sha256": canonical_payload,
    }


def validate_base(base: dict[str, Any]) -> dict[str, bool]:
    formal = base["paired_formal"]
    object_validation = base["object_validation"]
    f_certificate = formal["f"]["certificate"]
    return {
        "schema": base["schema"] == SCHEMA,
        "salt": base["construction_salt_sha256"] == SALT_SHA256,
        "initial_target_differ": base["initial_answer"] != base["target_answer"],
        "option_map_bijective": set(base["semantic_to_letter"].values())
        == set(OPTION_LETTERS),
        "object_solver_agreement": object_validation["passed"] is True,
        "object_single_f_error": object_validation["f_certificate"]["fatal_error_count"]
        == 1,
        "meta_v_solver_agreement": formal["v"]["agreement"] is True,
        "meta_v_unique_proof": formal["v"]["proof_count"] == 1,
        "meta_f_solver_agreement": formal["f"]["agreement"] is True,
        "meta_f_nonentailment": formal["f"]["python_label"] == "unknown"
        and formal["f"]["prover9_label"] == "unknown",
        "meta_f_countermodel": f_certificate["countermodel"]["model_found"] is True,
        "meta_f_exactly_one_error": f_certificate["fatal_error_count"] == 1,
        "meta_f_no_second_error": f_certificate["second_failure_point_absent"] is True,
        "v_f_edit_shared": base["pair_contract"]["proposed_edit_sha256"]
        == canonical_sha256(base["selected_edit"]),
        "v_f_output_contract_shared": base["pair_contract"]["output_contract"]
        == base["paired_formal"]["v"].get(
            "output_contract", base["pair_contract"]["output_contract"]
        ),
    }


def validate_arms(base: dict[str, Any], rows: list[dict[str, Any]]) -> dict[str, bool]:
    by_arm = {row["arm"]: row for row in rows}
    return {
        "four_arms": tuple(by_arm) == ARMS,
        "unique_ids": len({row["id"] for row in rows}) == 4,
        "shared_base": all(row["base_id"] == base["base_id"] for row in rows),
        "shared_prefix": len({object_sha256(row["messages"][:-1]) for row in rows}) == 1,
        "prompt_hashes": all(
            row["prompt_canonical_sha256"] == object_sha256(row["messages"])
            for row in rows
        ),
        "v_gold_target": by_arm["V"]["gold_answer"] == base["target_answer"],
        "non_v_gold_initial": all(
            by_arm[arm]["gold_answer"] == base["initial_answer"] for arm in ("N", "R", "F")
        ),
        "v_surface": by_arm["V"]["followup"] == base["english_v_text"],
        "f_surface": by_arm["F"]["followup"] == base["english_f_text"],
        "response_contract": all(
            row["messages"][-1]["content"].endswith(row["response_contract"])
            for row in rows
        ),
    }


def evidence_counts() -> dict[str, int]:
    return {
        "solver_traces": len(list((HERE / "solver_traces").glob("*.json"))),
        "countermodels": len(list((HERE / "countermodels").glob("*.json"))),
        "f_certificates": len(list((HERE / "f_certificates").glob("*.json"))),
    }


def solver_release_checks() -> dict[str, Any]:
    trace_paths = sorted((HERE / "solver_traces").glob("*.json"))
    counter_paths = sorted((HERE / "countermodels").glob("*.json"))
    prover_versions = Counter()
    mace_versions = Counter()
    for path in trace_paths:
        trace = read_json(path)["trace"]
        first = trace.splitlines()[0]
        prover_versions[first] += 1
    for path in counter_paths:
        payload = read_json(path)
        trace = payload.get("trace")
        if trace:
            mace_versions[trace.splitlines()[0]] += 1
    checks = {
        "all_prover_traces_dec_2007": set(prover_versions)
        == {"Prover9 (32) version Dec-2007, Dec 2007."},
        "all_mace_traces_dec_2007": set(mace_versions)
        == {"Mace4 (32) version Dec-2007, Dec 2007."},
    }
    return {
        "checks": checks,
        "prover_versions": dict(prover_versions),
        "mace_versions": dict(mace_versions),
        "parent_reference_binary_sha256": {
            "prover9_win32": PARENT_PROVER9_REFERENCE_SHA256,
            "mace4_win32": PARENT_MACE4_REFERENCE_SHA256,
        },
    }


def build_manifest(report: dict[str, Any]) -> dict[str, Any]:
    files = []
    for path in sorted(HERE.rglob("*")):
        if not path.is_file() or "__pycache__" in path.parts or path.name == "manifest.json":
            continue
        relative = path.relative_to(HERE).as_posix()
        item = {"path": relative, "raw_file_sha256": file_sha256(path)}
        if path.suffix == ".jsonl":
            item["canonical_content_sha256"] = canonical_jsonl_sha256(path)
        elif path.suffix == ".json":
            item["canonical_content_sha256"] = canonical_sha256(read_json(path))
        elif path.suffix in {".py", ".md"}:
            item["canonical_content_sha256"] = canonical_text_sha256(path)
        files.append(item)
    return {
        "schema": f"{SCHEMA}.manifest",
        "status": report["status"],
        "gpu_authorized": report["gpu_authorized"],
        "gpu_used": False,
        "model": {"repo": MODEL_REPO, "revision": MODEL_REVISION},
        "counts": report["counts"],
        "protocol_hashes": report["protocol_hashes"],
        "human_review_packet": report["human_review_packet"],
        "files": files,
    }


def validate(*, require_release_gate: bool = False) -> dict[str, Any]:
    protocols = protocol_hashes()
    bases = read_jsonl(HERE / "bases24.jsonl")
    arms = read_jsonl(HERE / "arms96.jsonl")
    audits = read_jsonl(HERE / "construction_audit24.jsonl")
    automatic = read_json(HERE / "automatic_validation_report.json")
    deterministic = read_json(HERE / "deterministic_rebuild_report.json")
    duplicates = read_json(HERE / "duplicate_audit.json")
    tokenizer = read_json(HERE / "tokenizer_preflight.json")
    packet = read_json(HERE / "HUMAN_REVIEW_PACKET.json")
    review_result = (
        read_json(HERE / "human_review_result.json")
        if (HERE / "human_review_result.json").is_file()
        else None
    )
    bases_by_id = {row["base_id"]: row for row in bases}
    arms_by_base: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in arms:
        arms_by_base[row["base_id"]].append(row)
    base_results = {base_id: validate_base(base) for base_id, base in bases_by_id.items()}
    arm_results = {
        base_id: validate_arms(base, arms_by_base[base_id])
        for base_id, base in bases_by_id.items()
    }
    family_variants = Counter((row["family_id"], row["variant"]) for row in bases)
    classes_by_family = Counter((row["error_class"], row["family_id"]) for row in bases)
    transitions = {
        variant: Counter(
            (row["initial_answer"], row["target_answer"])
            for row in bases
            if row["variant"] == variant
        )
        for variant in ("easy", "medium")
    }
    evidence = evidence_counts()
    solvers = solver_release_checks()
    workbook = validate_workbook(bases)
    checks = {
        "protocol_hashes_frozen": True,
        "automatic_construction_passed": automatic["status"] == "passed"
        and all(automatic["checks"].values()),
        "deterministic_two_builds": deterministic["status"] == "passed"
        and deterministic["two_clean_builds"] is True,
        "duplicate_audit_passed": duplicates["status"] == "passed"
        and all(duplicates["checks"].values()),
        "bases_24_unique": len(bases) == len(bases_by_id) == 24,
        "arms_96_unique": len(arms) == len({row["id"] for row in arms}) == 96,
        "families_12": len({row["family_id"] for row in bases}) == 12,
        "each_family_easy_medium": len(family_variants) == 24
        and set(family_variants.values()) == {1},
        "six_classes_two_families": set(row["error_class"] for row in bases)
        == set(ERROR_CLASSES)
        and len(classes_by_family) == 12
        and set(classes_by_family.values()) == {2},
        "each_variant_all_transitions": all(
            set(counter) == set(ORDERED_TRANSITIONS) and set(counter.values()) == {1}
            for counter in transitions.values()
        ),
        "base_invariants": all(all(result.values()) for result in base_results.values()),
        "arm_invariants": all(all(result.values()) for result in arm_results.values()),
        "stored_audits_passed": len(audits) == 24
        and all(row["passed_pre_tokenizer"] for row in audits),
        "evidence_files_complete": evidence
        == {"solver_traces": 192, "countermodels": 48, "f_certificates": 48},
        "solver_release_dec_2007": all(solvers["checks"].values()),
        "tokenizer_gate_passed": tokenizer["status"] == "passed"
        and all(tokenizer["checks"].values()),
        "workbook_blank_and_blind": workbook["status"] == "passed",
        "human_review_pending": packet["status"] == "awaiting_human_review"
        and packet["gpu_authorized"] is False,
        "construction_failures_retained": (HERE / "construction_failures.jsonl").is_file(),
        "no_qwen_outputs_used": automatic["qwen_outputs_used"] is False,
        "gpu_not_used": automatic["gpu_used"] is False and tokenizer["gpu_used"] is False,
    }
    release_checks = {
        "substantive_external_expert_review_passed": bool(
            review_result
            and review_result.get("status") == "passed_with_documentation_deviation"
            and review_result.get("substantive_human_review_passed") is True
            and review_result.get("pairs_reviewed") == 24
            and review_result.get("all_pairs_passed") is True
        ),
        "all_rubric_criteria_attested_passed": bool(
            review_result
            and review_result.get("all_prespecified_rubric_criteria_passed") is True
        ),
        "documentation_deviation_disclosed": bool(
            review_result
            and review_result.get("protocol_deviation", {}).get("present") is True
            and review_result.get("completed_workbook_available") is False
            and review_result.get("item_level_numeric_scores_available") is False
        ),
        "no_codex_fabricated_ratings": bool(
            review_result and review_result.get("codex_filled_human_ratings") is False
        ),
        "review_packet_hash_bound": bool(
            review_result
            and review_result.get("review_packet", {}).get("raw_sha256")
            == packet["workbook_raw_sha256"]
            and review_result.get("review_packet", {}).get("canonical_content_sha256")
            == packet["workbook_canonical_content_sha256"]
        ),
    }
    failed = [name for name, passed in checks.items() if not passed]
    if require_release_gate:
        failed.extend(
            f"release:{name}" for name, passed in release_checks.items() if not passed
        )
    gpu_authorized = require_release_gate and not failed
    report = {
        "schema": f"{SCHEMA}.independent_validation",
        "status": (
            "ready_for_gpu_confirmation"
            if gpu_authorized
            else "awaiting_human_review"
            if not failed
            else "failed"
        ),
        "passed_automatic_gate": not failed,
        "gpu_authorized": gpu_authorized,
        "gpu_used": False,
        "checks": checks,
        "release_checks": release_checks,
        "failures": failed,
        "counts": {
            "families": len({row["family_id"] for row in bases}),
            "bases": len(bases),
            "arms": len(arms),
            **evidence,
        },
        "protocol_hashes": protocols,
        "solver_release": solvers,
        "tokenizer": {
            "status": tokenizer["status"],
            "minimum_tokens": tokenizer["inputs"]["token_counts"]["minimum"],
            "maximum_tokens": tokenizer["inputs"]["token_counts"]["maximum"],
            "axis_sha256": tokenizer["official_p25"]["axis_sha256"],
            "capping_config_sha256": tokenizer["official_p25"]["capping_config_sha256"],
        },
        "workbook_validation": workbook,
        "human_review_packet": {
            "status": packet["status"],
            "workbook": packet["workbook"],
            "workbook_raw_sha256": packet["workbook_raw_sha256"],
            "workbook_canonical_content_sha256": packet[
                "workbook_canonical_content_sha256"
            ],
        },
        "base_results": base_results,
        "arm_results": arm_results,
    }
    write_json(HERE / "independent_validation_report.json", report)
    write_json(
        HERE / "STATUS.json",
        {
            "schema": f"{SCHEMA}.status",
            "status": report["status"],
            "automatic_gate_passed": report["passed_automatic_gate"],
            "human_review_required": not gpu_authorized,
            "gpu_authorized": gpu_authorized,
            "agent_may_fill_review": False,
        },
    )
    if gpu_authorized:
        write_json(
            HERE / "CPU_GATE.json",
            {
                "schema": f"{SCHEMA}.cpu_gate",
                "gate": "E2-new-salt-confirmation",
                "status": "passed_with_review_documentation_deviation",
                "gpu_authorized": True,
                "may_run_smoke": True,
                "formal_requires_passed_smoke_audit": True,
                "conditions": ["uncapped", "full_p25"],
                "smoke_rows": 48,
                "formal_rows": 192,
                "model": {"repo": MODEL_REPO, "revision": MODEL_REVISION},
                "p25_layers": list(range(46, 54)),
                "p1_authorized": False,
                "additional_salt_authorized": False,
                "data_hashes": {
                    "bases24_canonical_sha256": canonical_jsonl_sha256(
                        HERE / "bases24.jsonl"
                    ),
                    "arms96_canonical_sha256": canonical_jsonl_sha256(
                        HERE / "arms96.jsonl"
                    ),
                    "builder_raw_sha256": file_sha256(
                        HERE / "build_e2_new_salt_confirmation.py"
                    ),
                    "validator_raw_sha256": file_sha256(Path(__file__)),
                    "review_result_canonical_sha256": canonical_sha256(review_result),
                },
                "code_hashes": {
                    "run_common_raw_sha256": file_sha256(
                        HERE / "e2_new_salt_run_common.py"
                    ),
                    "runner_raw_sha256": file_sha256(
                        HERE.parents[2] / "scripts" / "run_e2_new_salt_confirmation.py"
                    ),
                    "auditor_raw_sha256": file_sha256(
                        HERE.parents[2] / "scripts" / "audit_e2_new_salt_confirmation.py"
                    ),
                    "analyzer_raw_sha256": file_sha256(
                        HERE.parents[2] / "scripts" / "analyze_e2_new_salt_confirmation.py"
                    ),
                },
                "protocol_hashes": protocols,
                "tokenizer_and_assets": {
                    "tokenizer_status": tokenizer["status"],
                    "axis_sha256": tokenizer["official_p25"]["axis_sha256"],
                    "capping_config_sha256": tokenizer["official_p25"][
                        "capping_config_sha256"
                    ],
                },
                "selection_fingerprints": {
                    scope: selection_fingerprint(scope)
                    for scope in ("smoke", "formal")
                },
                "deterministic_rebuild_passed": True,
                "duplicate_audit_passed": True,
                "human_review": {
                    "status": review_result["status"],
                    "pairs_reviewed": review_result["pairs_reviewed"],
                    "documentation_deviation": review_result["protocol_deviation"],
                    "codex_filled_human_ratings": False,
                },
                "gpu_used": False,
            },
        )
    write_json(HERE / "manifest.json", build_manifest(report))
    require(not failed, f"Independent automatic validation failed: {failed}")
    return report


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--require-release-gate", action="store_true")
    args = parser.parse_args()
    result = validate(require_release_gate=args.require_release_gate)
    print(
        json.dumps(
            {
                "status": result["status"],
                "counts": result["counts"],
                "workbook": result["human_review_packet"],
                "gpu_authorized": result["gpu_authorized"],
            },
            sort_keys=True,
        )
    )


if __name__ == "__main__":
    main()
