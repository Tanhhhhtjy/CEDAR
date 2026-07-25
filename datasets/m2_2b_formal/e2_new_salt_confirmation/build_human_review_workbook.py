#!/usr/bin/env python3
"""Create the blank, label-blind human-review workbook for all 24 pairs."""

from __future__ import annotations

import argparse
import hashlib
import json
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from e2_new_salt_common import (
    HERE,
    SALT,
    SCHEMA,
    file_sha256,
    object_sha256,
    require,
    write_json,
)
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

WORKBOOK_NAME = "e2_new_salt_blind_review24.xlsx"
SHEET_NAME = "盲审24"
FIRST_REVIEW_ROW = 6
LAST_REVIEW_ROW = 29


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line]


def blind_order(base_id: str) -> tuple[str, str]:
    digest = hashlib.sha256(f"{SALT}|blind-order|{base_id}".encode()).digest()
    return ("V", "F") if digest[0] % 2 == 0 else ("F", "V")


def review_id(base_id: str) -> str:
    digest = hashlib.sha256(f"{SALT}|review-id|{base_id}".encode()).hexdigest()
    return f"E2R-{digest[:10].upper()}"


def visible_payload(bases: list[dict[str, Any]]) -> dict[str, Any]:
    rows = []
    for base in bases:
        first, second = blind_order(base["base_id"])
        rows.append(
            {
                "review_id": review_id(base["base_id"]),
                "english_base_prompt": base["english_base_prompt"],
                "chinese_base_prompt": base["chinese_base_prompt"],
                "english_candidate_x": base[f"english_{first.lower()}_text"],
                "chinese_candidate_x": base[f"chinese_{first.lower()}_text"],
                "english_candidate_y": base[f"english_{second.lower()}_text"],
                "chinese_candidate_y": base[f"chinese_{second.lower()}_text"],
                "pair_fingerprint": object_sha256(
                    {
                        "base_id": base["base_id"],
                        "surface": {
                            key: base[key]
                            for key in (
                                "english_base_prompt",
                                "chinese_base_prompt",
                                "english_v_text",
                                "chinese_v_text",
                                "english_f_text",
                                "chinese_f_text",
                            )
                        },
                    }
                ),
            }
        )
    return {
        "schema": f"{SCHEMA}.blank_blind_review_workbook",
        "rows": rows,
        "ratings_blank": True,
        "gold_hidden": True,
        "f_class_hidden": True,
        "symbolic_internals_hidden": True,
        "qwen_outputs_absent": True,
    }


def write_workbook(path: Path, bases: list[dict[str, Any]]) -> str:
    payload = visible_payload(bases)
    packet_hash = object_sha256(payload)
    workbook = Workbook()
    workbook.properties.title = "CEDAR E2 new-salt blind human review"
    workbook.properties.keywords = packet_hash
    fixed_time = datetime(2026, 7, 25, tzinfo=UTC)
    workbook.properties.created = fixed_time
    workbook.properties.modified = fixed_time

    sheet = workbook.active
    sheet.title = SHEET_NAME
    guide = workbook.create_sheet("审核说明")
    attestation = workbook.create_sheet("审核声明")

    sheet.merge_cells("A1:Q1")
    sheet["A1"] = "CEDAR E2 new-salt 全量盲审（24 对）"
    sheet["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="274C77")
    sheet["A1"].alignment = Alignment(horizontal="center")
    sheet.merge_cells("A2:Q2")
    sheet["A2"] = "候选 X/Y 的条件身份已隐藏；请独立判断哪一项含有唯一致命推理错误。每格均须填写。"
    sheet.merge_cells("A3:Q3")
    sheet["A3"] = (
        "通过：X/Y 自然度均≥4；指代清楚；恰有一项含单一错误且选择正确；"
        "无答案/条件泄漏；风格匹配。"
    )
    headers = [
        "审核编号",
        "英文原题",
        "中文辅助",
        "英文候选X",
        "中文候选X",
        "英文候选Y",
        "中文候选Y",
        "X自然度1-5",
        "Y自然度1-5",
        "指代清晰",
        "含唯一错误的候选",
        "X答案泄漏",
        "Y答案泄漏",
        "条件身份泄漏",
        "X/Y风格匹配",
        "备注",
        "pair fingerprint（勿改）",
    ]
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=5, column=column, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F6D7A")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for offset, row in enumerate(payload["rows"], start=FIRST_REVIEW_ROW):
        values = [
            row["review_id"],
            row["english_base_prompt"],
            row["chinese_base_prompt"],
            row["english_candidate_x"],
            row["chinese_candidate_x"],
            row["english_candidate_y"],
            row["chinese_candidate_y"],
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            row["pair_fingerprint"],
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=offset, column=column, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.row_dimensions[offset].height = 180

    score = DataValidation(type="whole", operator="between", formula1="1", formula2="5")
    yes_no = DataValidation(type="list", formula1='"yes,no"')
    error_choice = DataValidation(type="list", formula1='"X,Y,none,both"')
    sheet.add_data_validation(score)
    sheet.add_data_validation(yes_no)
    sheet.add_data_validation(error_choice)
    score.add(f"H{FIRST_REVIEW_ROW}:I{LAST_REVIEW_ROW}")
    yes_no.add(f"J{FIRST_REVIEW_ROW}:J{LAST_REVIEW_ROW}")
    error_choice.add(f"K{FIRST_REVIEW_ROW}:K{LAST_REVIEW_ROW}")
    yes_no.add(f"L{FIRST_REVIEW_ROW}:O{LAST_REVIEW_ROW}")
    red = PatternFill("solid", fgColor="F4CCCC")
    sheet.conditional_formatting.add(
        f"H{FIRST_REVIEW_ROW}:I{LAST_REVIEW_ROW}",
        FormulaRule(formula=[f"H{FIRST_REVIEW_ROW}<4"], fill=red),
    )
    sheet.freeze_panes = f"A{FIRST_REVIEW_ROW}"
    widths = {
        "A": 17,
        "B": 54,
        "C": 48,
        "D": 54,
        "E": 48,
        "F": 54,
        "G": 48,
        "H": 12,
        "I": 12,
        "J": 12,
        "K": 18,
        "L": 12,
        "M": 12,
        "N": 16,
        "O": 15,
        "P": 32,
        "Q": 15,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.column_dimensions["Q"].hidden = True

    guide["A1"] = "审核说明"
    guide["A1"].font = Font(bold=True, size=14)
    lines = [
        "本表包含全部24对候选。X/Y顺序已盲化，表内不含gold、F类别、符号证明或Qwen输出。",
        "自然度：分别给X/Y打1--5分，二者都至少4分才通过。",
        "指代清晰：两项中的实体、记录和关系范围都清楚时填yes。",
        "含唯一错误的候选：根据论证选择X或Y；若都正确填none，若都有错误填both。",
        "答案泄漏：文字是否直接暴露原题答案；X和Y分别填写yes/no。",
        "条件身份泄漏：不做逻辑判断，仅凭固定语气、来源或格式能否识别哪项有错。",
        "X/Y风格匹配：两项是否像同一来源、同一言语行为和同一修改框架。",
        "任一行任一标准失败时整批需由builder完整重建并重新审核；不得只改失败题。",
    ]
    for row_index, line in enumerate(lines, start=3):
        guide.cell(row=row_index, column=1, value=line).alignment = Alignment(
            wrap_text=True, vertical="top"
        )
    guide.column_dimensions["A"].width = 120

    attestation["A1"] = "真实人工审核声明"
    attestation["A1"].font = Font(bold=True, size=14)
    attestation["A2"] = "审核者ID（内部匿名ID即可）"
    attestation["A3"] = "审核完成日期（YYYY-MM-DD）"
    attestation["A4"] = "本人是真实人工审核者（填yes）"
    attestation["A5"] = "签名或姓名缩写"
    for coordinate in ("B2", "B3", "B4", "B5"):
        attestation[coordinate] = None
    human_yes = DataValidation(type="list", formula1='"yes,no"')
    attestation.add_data_validation(human_yes)
    human_yes.add("B4")
    attestation.column_dimensions["A"].width = 42
    attestation.column_dimensions["B"].width = 36

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return packet_hash


def build(output_root: Path) -> dict[str, Any]:
    bases = read_jsonl(output_root / "bases24.jsonl")
    require(len(bases) == 24, "Blind review workbook requires exactly 24 bases")
    workbook_path = output_root / WORKBOOK_NAME
    packet_hash = write_workbook(workbook_path, bases)
    result = {
        "schema": f"{SCHEMA}.human_review_packet",
        "status": "awaiting_human_review",
        "review_mode": "single_real_human_full_review",
        "reviewed_pairs_required": 24,
        "workbook": WORKBOOK_NAME,
        "workbook_raw_sha256": file_sha256(workbook_path),
        "workbook_canonical_content_sha256": packet_hash,
        "blinding": {
            "gold_hidden": True,
            "f_class_hidden": True,
            "symbolic_internals_hidden": True,
            "qwen_outputs_absent": True,
            "candidate_order_deterministic": True,
        },
        "agent_filled_human_ratings": False,
        "gpu_authorized": False,
        "gpu_used": False,
    }
    write_json(output_root / "HUMAN_REVIEW_PACKET.json", result)
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-root", type=Path, default=HERE)
    args = parser.parse_args()
    result = build(args.output_root.resolve())
    print(json.dumps(result, sort_keys=True))


if __name__ == "__main__":
    main()
