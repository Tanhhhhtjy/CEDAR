"""Read-only import and Gate A/C finalization for a real v1.2 review workbook."""

from __future__ import annotations

import argparse
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from paired_justification_common import canonical_sha256, file_sha256, read_jsonl, write_json

HERE = Path(__file__).resolve().parent
REPO_ROOT = HERE.parents[2]
DEFAULT_TEMPLATE = HERE / "paired_justification_review12_bilingual.xlsx"
DEFAULT_RESULT = HERE / "human_review_result.json"
DEFAULT_GATE = REPO_ROOT / "explore_outputs" / "prover_generator_hardening_v1_2" / "FINAL_GATE.json"
REVIEW_SHEET = "自然化审核v1.2"
REVIEW_ROWS = range(6, 18)
ALLOWED_COLUMNS = set(range(8, 15))


class ReviewPacketError(ValueError):
    """Raised before outputs are changed when a completed workbook is invalid."""


def _verify_integrity(template_path: Path, completed_path: Path) -> dict[str, Any]:
    template = load_workbook(template_path, read_only=False, data_only=False)
    completed = load_workbook(completed_path, read_only=False, data_only=False)
    if completed.sheetnames != template.sheetnames:
        raise ReviewPacketError("sheet names or order changed")
    if completed.properties.keywords != template.properties.keywords:
        raise ReviewPacketError("packet metadata changed")
    checked = 0
    for name in template.sheetnames:
        expected = template[name]
        observed = completed[name]
        if (expected.max_row, expected.max_column) != (observed.max_row, observed.max_column):
            raise ReviewPacketError(f"sheet dimensions changed: {name}")
        for row in range(1, expected.max_row + 1):
            for column in range(1, expected.max_column + 1):
                if name == REVIEW_SHEET and row in REVIEW_ROWS and column in ALLOWED_COLUMNS:
                    continue
                checked += 1
                if expected.cell(row, column).value != observed.cell(row, column).value:
                    raise ReviewPacketError(
                        f"immutable cell changed: {name}!{observed.cell(row, column).coordinate}"
                    )
    return {
        "passed": True,
        "immutable_cells_checked": checked,
        "template_raw_sha256": file_sha256(template_path),
        "completed_raw_sha256": file_sha256(completed_path),
        "packet_metadata": template.properties.keywords,
    }


def _yes_no(value: Any, coordinate: str) -> str:
    if not isinstance(value, str) or value.strip().lower() not in {"yes", "no"}:
        raise ReviewPacketError(f"{coordinate} must be yes or no")
    return value.strip().lower()


def _score(value: Any, coordinate: str) -> int:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ReviewPacketError(f"{coordinate} must be an integer from 1 to 5")
    score = int(value)
    if value != score or not 1 <= score <= 5:
        raise ReviewPacketError(f"{coordinate} must be an integer from 1 to 5")
    return score


def finalize(
    completed_path: Path,
    *,
    template_path: Path = DEFAULT_TEMPLATE,
    result_path: Path = DEFAULT_RESULT,
    gate_path: Path = DEFAULT_GATE,
) -> dict[str, Any]:
    integrity = _verify_integrity(template_path.resolve(), completed_path.resolve())
    source = read_jsonl(HERE / "paired_justification_sentinels12.jsonl")
    source_by_id = {row["pair_id"]: row for row in source}
    workbook = load_workbook(completed_path, read_only=True, data_only=False)
    sheet = workbook[REVIEW_SHEET]
    rows = []
    for row_number in REVIEW_ROWS:
        review_id = sheet.cell(row_number, 1).value
        if review_id not in source_by_id:
            raise ReviewPacketError(f"unknown review id at row {row_number}")
        source_hash = sheet.cell(row_number, 15).value
        if source_hash != source_by_id[review_id]["source_symbolic_sha256"]:
            raise ReviewPacketError(f"source hash changed for {review_id}")
        naturalness = _score(sheet.cell(row_number, 8).value, f"H{row_number}")
        clear_reference = _yes_no(sheet.cell(row_number, 9).value, f"I{row_number}")
        single_error = _yes_no(sheet.cell(row_number, 10).value, f"J{row_number}")
        gold_leakage = _yes_no(sheet.cell(row_number, 11).value, f"K{row_number}")
        condition_leakage = _yes_no(sheet.cell(row_number, 12).value, f"L{row_number}")
        style_matched = _yes_no(sheet.cell(row_number, 13).value, f"M{row_number}")
        notes = sheet.cell(row_number, 14).value
        if notes is not None and not isinstance(notes, str):
            raise ReviewPacketError(f"N{row_number} must be text or blank")
        checks = {
            "language_naturalness_gte_4": naturalness >= 4,
            "clear_reference_yes": clear_reference == "yes",
            "single_interpretable_error_yes": single_error == "yes",
            "surface_gold_leakage_no": gold_leakage == "no",
            "condition_identity_leakage_no": condition_leakage == "no",
            "v_f_style_matched_yes": style_matched == "yes",
        }
        rows.append(
            {
                "review_id": review_id,
                "source_symbolic_sha256": source_hash,
                "language_naturalness": naturalness,
                "clear_reference": clear_reference,
                "single_interpretable_error": single_error,
                "surface_gold_leakage": gold_leakage,
                "condition_identity_leakage": condition_leakage,
                "v_f_style_matched": style_matched,
                "notes": notes,
                "checks": checks,
                "passed": all(checks.values()),
            }
        )
    all_passed = all(row["passed"] for row in rows)
    result = {
        "schema": "cedar.prover_generator_hardening.v1_2.human_review_result",
        "status": "passed" if all_passed else "failed",
        "reviewed_row_count": 12,
        "complete": True,
        "integrity": integrity,
        "item_level_all_must_pass": True,
        "passed_row_count": sum(row["passed"] for row in rows),
        "failed_review_ids": [row["review_id"] for row in rows if not row["passed"]],
        "rows": rows,
        "result_canonical_sha256": canonical_sha256(rows),
        "agent_filled_human_ratings": False,
    }
    gate = {
        "gate": "A" if all_passed else "C",
        "status": "human_review_passed" if all_passed else "naturalization_failed",
        "step_1_complete": all_passed,
        "may_prepare_development_bank_protocol": all_passed,
        "renderer_revision_only": not all_passed,
        "human_review_result_canonical_sha256": canonical_sha256(result),
        "gpu_qwen_capping_run": False,
    }
    if not all_passed:
        gate["failed_review_ids"] = result["failed_review_ids"]
    write_json(result_path.resolve(), result)
    write_json(gate_path.resolve(), gate)
    return {"result": result, "gate": gate}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("completed_workbook", type=Path)
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--result", type=Path, default=DEFAULT_RESULT)
    parser.add_argument("--gate", type=Path, default=DEFAULT_GATE)
    args = parser.parse_args()
    try:
        output = finalize(
            args.completed_workbook,
            template_path=args.template,
            result_path=args.result,
            gate_path=args.gate,
        )
    except ReviewPacketError as error:
        print({"status": "invalid_review_packet", "error": str(error)})
        raise SystemExit(2) from error
    print(
        {
            "status": output["result"]["status"],
            "gate": output["gate"]["gate"],
            "failed_review_ids": output["result"]["failed_review_ids"],
        }
    )


if __name__ == "__main__":
    main()
