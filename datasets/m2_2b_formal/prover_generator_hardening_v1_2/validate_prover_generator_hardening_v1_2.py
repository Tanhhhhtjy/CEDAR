# ruff: noqa: E501
"""Independent hard-gate validation for the twelve v1.2 paired sentinels."""

from __future__ import annotations

import argparse
import copy
import json
from collections import Counter
from pathlib import Path
from typing import Any

from build_prover_generator_hardening_v1_2 import WORKBOOK_NAME, validate_row
from openpyxl import load_workbook
from paired_justification_common import (
    ERROR_CLASSES,
    V1_1,
    canonical_json_file_sha256,
    canonical_jsonl_sha256,
    canonical_sha256,
    derive_error_class,
    evaluate_fixed_point,
    fact,
    literal_key,
    proof_count,
    read_jsonl,
    run_mace4_countermodel,
    shortest_proof,
    validate_with_both,
    write_json,
)

HERE = Path(__file__).resolve().parent


def validate_formal(row: dict[str, Any]) -> dict[str, Any]:
    formal = row["formal"]
    query = formal["query"]
    v_world = formal["v"]["world"]
    f_world = formal["f"]["world"]
    certificate = formal["f"]["certificate"]
    v_dual = validate_with_both(v_world, query)
    f_dual = validate_with_both(f_world, query)
    v_proof = shortest_proof(v_dual["fixed_point"], query)
    f_claim = formal["f"]["claim"]
    f_premises = formal["f"]["premises"]
    derived = derive_error_class(f_world, f_premises, f_claim)
    assumed_world = copy.deepcopy(f_world)
    assumed_world["records"].append(
        fact("independent_temporary_assumption", f_claim, role="fallacy_assumption")
    )
    assumed = evaluate_fixed_point(assumed_world)
    countermodel = run_mace4_countermodel(f_world, query)
    checks = {
        "v_dual_true": v_dual["agreement"]
        and v_dual["python_label"] == v_dual["prover9_label"] == "true",
        "v_unique_proof": proof_count(v_dual["fixed_point"], query) == 1,
        "v_proof_matches_stored": canonical_sha256(v_proof)
        == canonical_sha256(formal["v"]["proof"]),
        "f_dual_unknown": f_dual["agreement"]
        and f_dual["python_label"] == f_dual["prover9_label"] == "unknown",
        "f_class_derived_independently": derived
        == row["error_class"]
        == certificate["declared_error_class"],
        "f_premises_entailed": all(
            literal_key(item) in f_dual["fixed_point"]["closure"] for item in f_premises
        ),
        "f_claim_not_entailed": literal_key(f_claim) not in f_dual["fixed_point"]["closure"],
        "f_downstream_valid_with_one_assumption": all(
            literal_key(item) in assumed["closure"] for item in certificate["downstream_literals"]
        ),
        "f_warrant_absent_without_assumption": literal_key(query)
        not in f_dual["fixed_point"]["closure"],
        "f_warrant_present_with_assumption": literal_key(query) in assumed["closure"],
        "f_countermodel": countermodel["model_found"],
        "f_fatal_error_count_one": certificate["fatal_error_count"] == 1,
        "f_no_second_error": certificate["second_failure_point_absent"] is True,
    }
    return {"passed": all(checks.values()), "checks": checks}


def validate_workbook(root: Path, rows: list[dict[str, Any]]) -> dict[str, Any]:
    path = root / WORKBOOK_NAME
    workbook = load_workbook(path, read_only=False, data_only=False)
    checks: dict[str, bool] = {
        "expected_sheets": workbook.sheetnames == ["自然化审核v1.2", "审核说明"],
        "packet_hash_bound": workbook.properties.keywords
        == json.loads((root / "manifest.json").read_text(encoding="utf-8"))[
            "workbook_canonical_content_sha256"
        ],
    }
    sheet = workbook["自然化审核v1.2"]
    checks["twelve_rows"] = sheet.max_row == 17
    checks["all_ratings_blank"] = all(
        sheet.cell(row=row, column=column).value is None
        for row in range(6, 18)
        for column in range(8, 15)
    )
    checks["ids_match"] = [sheet.cell(row=row, column=1).value for row in range(6, 18)] == [
        item["pair_id"] for item in rows
    ]
    visible_surface = "\n".join(
        str(sheet.cell(row=row, column=column).value or "")
        for row in range(1, sheet.max_row + 1)
        for column in range(1, 15)
    ).lower()
    checks["internals_hidden"] = all(
        marker not in visible_surface
        for marker in [*ERROR_CLASSES, "python_label", "prover9_label", "revision_warranted"]
    )
    checks["source_hash_column_hidden"] = sheet.column_dimensions["O"].hidden is True
    return {"passed": all(checks.values()), "checks": checks}


def validate(root: Path) -> dict[str, Any]:
    rows = read_jsonl(root / "paired_justification_sentinels12.jsonl")
    source_by_id = {
        row["attempt_id"]: row for row in read_jsonl(V1_1 / "hybrid_stress48.jsonl")
    }
    pair_results: dict[str, Any] = {}
    for row in rows:
        recomputed_surface = validate_row(row)
        formal = validate_formal(row)
        source = source_by_id[row["source_attempt_id"]]
        source_checks = {
            "source_symbolic_hash": row["source_symbolic_sha256"] == canonical_sha256(source),
            "source_edit_hash": row["source_selected_edit_sha256"]
            == canonical_sha256(source["validation"]["selected_edit"]),
            "surface_edit_unchanged": canonical_sha256(row["surface"]["proposed_edit"])
            == canonical_sha256(source["validation"]["selected_edit"]),
        }
        pair_results[row["pair_id"]] = {
            "surface": recomputed_surface,
            "formal": formal,
            "source": {"passed": all(source_checks.values()), "checks": source_checks},
            "passed": recomputed_surface["passed"]
            and formal["passed"]
            and all(source_checks.values()),
        }
    matrix = Counter((row["error_class"], row["variant"]) for row in rows)
    workbook = validate_workbook(root, rows)
    determinism = json.loads(
        (root / "deterministic_rebuild_report.json").read_text(encoding="utf-8")
    )
    automatic = json.loads(
        (root / "automatic_validation_report.json").read_text(encoding="utf-8")
    )
    manifest = json.loads((root / "manifest.json").read_text(encoding="utf-8"))
    top_checks = {
        "exactly_12": len(rows) == 12,
        "matrix_complete": matrix
        == Counter({(name, variant): 1 for name in ERROR_CLASSES for variant in ("easy", "medium")}),
        "all_pairs": all(item["passed"] for item in pair_results.values()),
        "automatic_report_passed": automatic["passed"],
        "determinism_passed": determinism["passed"],
        "workbook_passed": workbook["passed"],
        "manifest_sentinel_hash": canonical_jsonl_sha256(
            root / "paired_justification_sentinels12.jsonl"
        )
        == next(
            item["canonical_content_sha256"]
            for item in manifest["files"]
            if item["path"].endswith("paired_justification_sentinels12.jsonl")
        ),
        "manifest_automatic_hash": canonical_json_file_sha256(
            root / "automatic_validation_report.json"
        )
        == next(
            item["canonical_content_sha256"]
            for item in manifest["files"]
            if item["path"].endswith("automatic_validation_report.json")
        ),
        "no_model_execution": manifest["gpu_qwen_capping_run"] is False,
        "no_development_bank": manifest["development_bank_built"] is False,
    }
    return {
        "schema": "cedar.prover_generator_hardening.v1_2.independent_validation",
        "passed": all(top_checks.values()),
        "checks": top_checks,
        "pair_results": pair_results,
        "workbook": workbook,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=HERE)
    args = parser.parse_args()
    report = validate(args.root.resolve())
    write_json(args.root.resolve() / "independent_validation_report.json", report)
    print({"passed": report["passed"], "pairs": len(report["pair_results"])})
    if not report["passed"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
