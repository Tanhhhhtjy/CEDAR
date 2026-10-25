# ruff: noqa: E501
"""Build exactly twelve v1.2 paired-justification sentinels and a blank review packet."""

from __future__ import annotations

import argparse
import copy
import json
import re
import tempfile
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from paired_justification_common import (
    AUTHENTICATION_STATUS,
    ERROR_CLASSES,
    FORBIDDEN_PHRASES,
    MASTER_SEED,
    REVIEWER,
    SCHEMA,
    SOURCE_COMMIT,
    SOURCE_TYPE,
    SPEECH_ACT,
    V1_1,
    build_formal_pair,
    canonical_json_file_sha256,
    canonical_jsonl_sha256,
    canonical_sha256,
    file_sha256,
    render_pair,
    select_source_rows,
    stable_text_sha256,
    write_json,
    write_jsonl,
)

HERE = Path(__file__).resolve().parent
REPO_ROOT = HERE.parents[2]
EXPLORE = REPO_ROOT / "explore_outputs" / "prover_generator_hardening_v1_2"
WORKBOOK_NAME = "paired_justification_review12_bilingual.xlsx"


def _token_ngrams(text: str, size: int) -> set[str]:
    tokens = re.findall(r"[a-z]+", text.lower())
    return {" ".join(tokens[index : index + size]) for index in range(len(tokens) - size + 1)}


def _exclusive_fixed_phrases(rows: list[dict[str, Any]]) -> dict[str, Any]:
    report: dict[str, Any] = {}
    for size in range(2, 7):
        common_v = set.intersection(*(_token_ngrams(row["surface"]["english_v_text"], size) for row in rows))
        common_f = set.intersection(*(_token_ngrams(row["surface"]["english_f_text"], size) for row in rows))
        report[str(size)] = {
            "v_only": sorted(common_v - common_f),
            "f_only": sorted(common_f - common_v),
        }
    return report


def _coreference_checks(text: str) -> dict[str, bool]:
    lowered = text.lower()
    return {
        "no_repeated_indefinite_item": lowered.count("an item") <= 1,
        "unary_rule_uses_same_item": "if an item" not in lowered or "that same item" in lowered,
        "binary_source_named": "route" not in lowered or "source k" in lowered,
        "binary_target_named": "route" not in lowered or "target r" in lowered,
    }


def validate_row(row: dict[str, Any]) -> dict[str, Any]:
    formal = row["formal"]
    surface = row["surface"]
    v_contract = row["pair_contract"]["v"]
    f_contract = row["pair_contract"]["f"]
    ratio = len(surface["english_v_text"].split()) / len(surface["english_f_text"].split())
    depth = formal["v"]["proof"]["depth"]
    low, high = (2, 3) if row["variant"] == "easy" else (4, 6)
    certificate = formal["f"]["certificate"]
    all_surface = " ".join(
        [
            surface["english_v_text"],
            surface["english_f_text"],
        ]
    ).lower()
    coreference = {
        "v": _coreference_checks(surface["english_v_text"]),
        "f": _coreference_checks(surface["english_f_text"]),
    }
    checks = {
        "paired_slot_skeleton_hash_equal": v_contract["slot_skeleton_sha256"]
        == f_contract["slot_skeleton_sha256"],
        "reviewer_equal": v_contract["reviewer"] == f_contract["reviewer"] == REVIEWER,
        "source_type_equal": v_contract["source_type"] == f_contract["source_type"] == SOURCE_TYPE,
        "authentication_equal": v_contract["authentication_status"]
        == f_contract["authentication_status"]
        == AUTHENTICATION_STATUS,
        "speech_act_equal": v_contract["speech_act"] == f_contract["speech_act"] == SPEECH_ACT,
        "proposed_edit_hash_equal": v_contract["proposed_edit_sha256"]
        == f_contract["proposed_edit_sha256"],
        "proposed_edit_text_en_equal": v_contract["proposed_edit_text_en"]
        == f_contract["proposed_edit_text_en"],
        "proposed_edit_text_zh_equal": v_contract["proposed_edit_text_zh"]
        == f_contract["proposed_edit_text_zh"],
        "v_python_proves_warrant": formal["v"]["python_label"] == "true",
        "v_prover9_proves_warrant": formal["v"]["prover9_label"] == "true",
        "v_dual_solver_agreement": formal["v"]["agreement"],
        "v_unique_proof": formal["v"]["proof_count"] == 1,
        "v_depth_in_contract": low <= depth <= high,
        "f_python_does_not_prove_warrant": formal["f"]["python_label"] == "unknown",
        "f_prover9_does_not_prove_warrant": formal["f"]["prover9_label"] == "unknown",
        "f_dual_solver_agreement": formal["f"]["agreement"],
        "f_countermodel": certificate["countermodel"]["model_found"] is True,
        "f_declared_class_derived": certificate["declared_error_class"]
        == certificate["derived_error_class"]
        == row["error_class"],
        "f_fatal_error_count_one": certificate["fatal_error_count"] == 1,
        "f_claim_not_entailed": certificate["claim_not_entailed"],
        "f_downstream_valid_under_single_assumption": certificate[
            "downstream_valid_under_single_assumption"
        ],
        "f_warrant_absent_without_error": certificate["warrant_absent_without_invalid_step"],
        "f_no_second_error": certificate["second_failure_point_absent"],
        "length_ratio": 0.90 <= ratio <= 1.10,
        "coreference": all(all(values.values()) for values in coreference.values()),
        "forbidden_language_absent": all(phrase.lower() not in all_surface for phrase in FORBIDDEN_PHRASES),
        "no_encoding_artifacts": not any(
            marker in " ".join(
                [
                    surface["english_v_text"],
                    surface["english_f_text"],
                    surface["chinese_v_text"],
                    surface["chinese_f_text"],
                    surface["english_base_prompt"],
                    surface["chinese_base_prompt"],
                ]
            )
            for marker in ("Ą", "鈥", "�", "锛", "銆")
        ),
        "source_edit_unchanged": canonical_sha256(surface["proposed_edit"])
        == row["source_selected_edit_sha256"],
    }
    return {
        "passed": all(checks.values()),
        "checks": checks,
        "v_f_word_count_ratio": ratio,
        "v_word_count": len(surface["english_v_text"].split()),
        "f_word_count": len(surface["english_f_text"].split()),
        "v_warrant_proof_depth": depth,
        "coreference_checks": coreference,
    }


def _row(source: dict[str, Any], index: int) -> dict[str, Any]:
    error_class = source["assigned_f_error_class"]
    variant = source["variant"]
    edit_id = f"edit{index:02d}"
    formal = build_formal_pair(error_class, variant, edit_id)
    surface = render_pair(source, index, formal)
    shared_contract = {
        "reviewer": REVIEWER,
        "source_type": SOURCE_TYPE,
        "authentication_status": AUTHENTICATION_STATUS,
        "speech_act": SPEECH_ACT,
        "slot_skeleton_sha256": surface["slot_skeleton_sha256"],
        "proposed_edit_sha256": surface["proposed_edit_sha256"],
        "proposed_edit_text_en": surface["proposed_edit_text_en"],
        "proposed_edit_text_zh": surface["proposed_edit_text_zh"],
    }
    row = {
        "schema": SCHEMA,
        "pair_id": f"PJ-{index:02d}",
        "sentinel": True,
        "error_class": error_class,
        "variant": variant,
        "source_commit": SOURCE_COMMIT,
        "source_attempt_id": source["attempt_id"],
        "source_symbolic_sha256": canonical_sha256(source),
        "source_selected_edit_sha256": canonical_sha256(source["validation"]["selected_edit"]),
        "pair_contract": {"v": copy.deepcopy(shared_contract), "f": copy.deepcopy(shared_contract)},
        "formal": formal,
        "surface": surface,
        "gpu_qwen_capping_run": False,
    }
    row["automatic_validation"] = validate_row(row)
    return row


def _automatic_report(rows: list[dict[str, Any]]) -> dict[str, Any]:
    class_variant = Counter((row["error_class"], row["variant"]) for row in rows)
    pair_results = {row["pair_id"]: row["automatic_validation"] for row in rows}
    gradients = {}
    for error_class in ERROR_CLASSES:
        easy = next(row for row in rows if row["error_class"] == error_class and row["variant"] == "easy")
        medium = next(row for row in rows if row["error_class"] == error_class and row["variant"] == "medium")
        gradients[error_class] = {
            "easy": easy["automatic_validation"]["v_warrant_proof_depth"],
            "medium": medium["automatic_validation"]["v_warrant_proof_depth"],
            "passed": medium["automatic_validation"]["v_warrant_proof_depth"]
            > easy["automatic_validation"]["v_warrant_proof_depth"],
        }
    exclusive = _exclusive_fixed_phrases(rows)
    exclusive_empty = all(
        not values["v_only"] and not values["f_only"] for values in exclusive.values()
    )
    checks = {
        "exactly_12_pairs": len(rows) == 12,
        "class_variant_matrix": class_variant
        == Counter({(name, variant): 1 for name in ERROR_CLASSES for variant in ("easy", "medium")}),
        "all_row_gates": all(result["passed"] for result in pair_results.values()),
        "all_difficulty_gradients": all(item["passed"] for item in gradients.values()),
        "no_global_condition_exclusive_2_to_6_grams": exclusive_empty,
    }
    return {
        "schema": f"{SCHEMA}.automatic_checks",
        "passed": all(checks.values()),
        "checks": checks,
        "pair_results": pair_results,
        "difficulty_gradients": gradients,
        "condition_exclusive_fixed_phrases": exclusive,
        "template_distribution": {
            "template": Counter(row["surface"]["slot_skeleton"]["template"] for row in rows),
            "by_error_class": Counter(row["error_class"] for row in rows),
            "by_difficulty": Counter(row["variant"] for row in rows),
        },
        "gpu_qwen_capping_run": False,
    }


def _workbook_payload(rows: list[dict[str, Any]]) -> dict[str, Any]:
    fields = (
        "review_id",
        "english_base_prompt",
        "chinese_base_prompt",
        "english_v_text",
        "chinese_v_text",
        "english_f_text",
        "chinese_f_text",
        "source_symbolic_sha256",
    )
    return {
        "schema": f"{SCHEMA}.blank_review_workbook",
        "rows": [{field: row["surface"][field] for field in fields} for row in rows],
        "ratings_blank": True,
    }


def write_workbook(path: Path, rows: list[dict[str, Any]]) -> str:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "自然化审核v1.2"
    guide = workbook.create_sheet("审核说明")
    packet_hash = canonical_sha256(_workbook_payload(rows))
    workbook.properties.title = "CEDAR v1.2 配对论证自然化审核"
    workbook.properties.keywords = packet_hash

    sheet.merge_cells("A1:O1")
    sheet["A1"] = "CEDAR v1.2 配对论证自然化审核（空白人工表）"
    sheet["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="274C77")
    sheet["A1"].alignment = Alignment(horizontal="center")
    sheet.merge_cells("A2:O2")
    sheet["A2"] = "请只评价语言与构念隔离；不要推测隐藏标签。每一项都必须填写。"
    sheet.merge_cells("A3:O3")
    sheet["A3"] = "通过标准：自然度≥4；指代清晰=是；单一可解释错误=是；答案泄漏=否；条件身份泄漏=否；V/F风格匹配=是。"
    headers = [
        "编号",
        "英文原题",
        "中文辅助",
        "英文V",
        "中文V",
        "英文F",
        "中文F",
        "自然度1-5",
        "指代清晰",
        "单一可解释错误",
        "答案泄漏",
        "条件身份泄漏",
        "V/F风格匹配",
        "备注",
        "source hash（勿改）",
    ]
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=5, column=column, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F6D7A")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for offset, row in enumerate(rows, start=6):
        surface = row["surface"]
        values = [
            surface["review_id"],
            surface["english_base_prompt"],
            surface["chinese_base_prompt"],
            surface["english_v_text"],
            surface["chinese_v_text"],
            surface["english_f_text"],
            surface["chinese_f_text"],
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            surface["source_symbolic_sha256"],
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=offset, column=column, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.row_dimensions[offset].height = 165
    score_validation = DataValidation(type="whole", operator="between", formula1="1", formula2="5")
    yes_no_validation = DataValidation(type="list", formula1='"yes,no"')
    sheet.add_data_validation(score_validation)
    sheet.add_data_validation(yes_no_validation)
    score_validation.add("H6:H17")
    yes_no_validation.add("I6:M17")
    red = PatternFill("solid", fgColor="F4CCCC")
    sheet.conditional_formatting.add("H6:H17", FormulaRule(formula=["H6<4"], fill=red))
    sheet.freeze_panes = "A6"
    widths = {"A": 11, "B": 54, "C": 48, "D": 54, "E": 48, "F": 54, "G": 48, "H": 12, "I": 12, "J": 17, "K": 12, "L": 16, "M": 15, "N": 32, "O": 15}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.column_dimensions["O"].hidden = True

    guide["A1"] = "审核说明"
    guide["A1"].font = Font(bold=True, size=14)
    guide["A3"] = "condition_identity_leakage：不分析具体逻辑，仅凭固定来源、身份、语气或言语行为，是否能轻易判断V/F。"
    guide["A4"] = "single_interpretable_error：F是否只有一个可解释的致命推理错误；不要查看隐藏的错误类型。"
    guide["A5"] = "若任一行任一标准失败，该轮为自然化Gate C；不得以平均分抵消。"
    guide["A6"] = "本表没有gold、错误类别或符号结构。请勿修改题干、V/F文字、编号或隐藏source hash。"
    guide.column_dimensions["A"].width = 120
    for row in range(1, 7):
        guide.cell(row=row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return packet_hash


def _write_guide(path: Path) -> None:
    path.write_text(
        """# CEDAR v1.2 自然化人工审核指南

本表只评估自然语言质量和V/F构念隔离，不要求审核符号证明。每行必须填写六项。

- `language_naturalness`：1--5，至少4才通过。
- `clear_reference`：实体、同一变量和关系范围是否清楚。
- `single_interpretable_error`：F是否只有一个可定位的推理错误。
- `surface_gold_leakage`：文字是否直接提示原题三选一答案。
- `condition_identity_leakage`：不分析逻辑，仅凭来源、身份、语气或言语行为能否识别V/F。
- `v_f_style_matched`：两臂是否像同一复核者用同一渠道提出同类修改。

任一行不满足全部六项即失败。请勿修改题干、编号、V/F文本或隐藏source hash；只填写评分列和备注。
""",
        encoding="utf-8",
    )


def build(output_root: Path) -> dict[str, Any]:
    output_root.mkdir(parents=True, exist_ok=True)
    rows = [_row(source, index) for index, source in enumerate(select_source_rows(), start=1)]
    write_jsonl(output_root / "paired_justification_sentinels12.jsonl", rows)
    report = _automatic_report(rows)
    write_json(output_root / "automatic_validation_report.json", report)
    workbook_hash = write_workbook(output_root / WORKBOOK_NAME, rows)
    _write_guide(output_root / "NATURALIZATION_REVIEW_GUIDE_ZH.md")
    for row in rows:
        pair_id = row["pair_id"]
        write_json(output_root / "solver_traces" / f"{pair_id}__v_prover9.json", row["formal"]["v"]["solver"])
        write_json(output_root / "solver_traces" / f"{pair_id}__f_prover9.json", row["formal"]["f"]["solver"])
        write_json(output_root / "countermodels" / f"{pair_id}__f_countermodel.json", row["formal"]["f"]["certificate"]["countermodel"])
        write_json(output_root / "semantic_error_certificates" / f"{pair_id}__f_certificate.json", row["formal"]["f"]["certificate"])
    return {
        "rows": rows,
        "report": report,
        "workbook_canonical_content_sha256": workbook_hash,
        "sentinels_canonical_sha256": canonical_jsonl_sha256(output_root / "paired_justification_sentinels12.jsonl"),
        "automatic_report_canonical_sha256": canonical_json_file_sha256(output_root / "automatic_validation_report.json"),
    }


def verify_determinism(output_root: Path) -> dict[str, Any]:
    with tempfile.TemporaryDirectory() as first_temp, tempfile.TemporaryDirectory() as second_temp:
        first = build(Path(first_temp))
        second = build(Path(second_temp))
    compared = {
        "sentinels": first["sentinels_canonical_sha256"] == second["sentinels_canonical_sha256"],
        "automatic_report": first["automatic_report_canonical_sha256"] == second["automatic_report_canonical_sha256"],
        "workbook_content": first["workbook_canonical_content_sha256"] == second["workbook_canonical_content_sha256"],
    }
    result = {
        "schema": f"{SCHEMA}.determinism",
        "passed": all(compared.values()),
        "compared": compared,
        "first": {key: value for key, value in first.items() if key not in {"rows", "report"}},
        "second": {key: value for key, value in second.items() if key not in {"rows", "report"}},
    }
    write_json(output_root / "deterministic_rebuild_report.json", result)
    return result


def _manifest(output_root: Path, build_result: dict[str, Any], determinism: dict[str, Any]) -> dict[str, Any]:
    files = []
    workbook_hash = build_result["workbook_canonical_content_sha256"]
    for included_root in (output_root, EXPLORE):
        for path in sorted(included_root.rglob("*")):
            if not path.is_file() or path.name == "manifest.json" or "__pycache__" in path.parts:
                continue
            if path.suffix == ".jsonl":
                canonical_hash = canonical_jsonl_sha256(path)
            elif path.suffix == ".json":
                canonical_hash = canonical_json_file_sha256(path)
            elif path.suffix == ".xlsx":
                canonical_hash = workbook_hash
            else:
                canonical_hash = stable_text_sha256(path.read_text(encoding="utf-8"))
            files.append(
                {
                    "path": path.relative_to(REPO_ROOT).as_posix(),
                    "raw_file_sha256": file_sha256(path),
                    "canonical_content_sha256": canonical_hash,
                }
            )
    solver = json.loads((V1_1 / "f_sentinel_freeze.json").read_text(encoding="utf-8"))["solver"]
    return {
        "schema": f"{SCHEMA}.manifest",
        "research_anchor": SOURCE_COMMIT,
        "protocol_path": "docs/PROVER_GENERATOR_HARDENING_V1_2_PROTOCOL.md",
        "master_seed": MASTER_SEED,
        "solver": solver,
        "builder_sha256": file_sha256(output_root / "build_prover_generator_hardening_v1_2.py"),
        "validator_sha256": file_sha256(output_root / "validate_prover_generator_hardening_v1_2.py"),
        "formal_renderer_common_sha256": file_sha256(output_root / "paired_justification_common.py"),
        "human_review_finalizer_sha256": file_sha256(output_root / "finalize_paired_justification_review.py"),
        "protocol_sha256": file_sha256(
            REPO_ROOT / "docs" / "PROVER_GENERATOR_HARDENING_V1_2_PROTOCOL.md"
        ),
        "v1_1_source_jsonl_canonical_sha256": canonical_jsonl_sha256(V1_1 / "hybrid_stress48.jsonl"),
        "sentinel_count": 12,
        "paired_justification_only": True,
        "automatic_gate_passed": build_result["report"]["passed"],
        "deterministic_rebuild_passed": determinism["passed"],
        "human_review_status": "awaiting_human_naturalization_review",
        "workbook_canonical_content_sha256": workbook_hash,
        "files": files,
        "gpu_qwen_capping_run": False,
        "development_bank_built": False,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-root", type=Path, default=HERE)
    parser.add_argument("--skip-determinism", action="store_true")
    args = parser.parse_args()
    output_root = args.output_root.resolve()
    result = build(output_root)
    determinism = (
        {"passed": False, "skipped": True}
        if args.skip_determinism
        else verify_determinism(output_root)
    )
    gate = {
        "gate": "B" if result["report"]["passed"] and determinism["passed"] else "D",
        "status": "awaiting_human_naturalization_review"
        if result["report"]["passed"] and determinism["passed"]
        else "paired_justification_automatic_gate_failed",
        "sentinel_count": 12,
        "step_1_complete": False,
        "may_prepare_development_bank_protocol": False,
        "review_packet": WORKBOOK_NAME,
        "gpu_qwen_capping_run": False,
    }
    EXPLORE.mkdir(parents=True, exist_ok=True)
    write_json(EXPLORE / "FINAL_GATE.json", gate)
    manifest = _manifest(output_root, result, determinism)
    write_json(output_root / "manifest.json", manifest)
    print(
        {
            "automatic_gate": result["report"]["passed"],
            "deterministic_rebuild": determinism["passed"],
            "final_gate": gate["gate"],
            "human_review_status": gate["status"],
        }
    )


if __name__ == "__main__":
    main()
