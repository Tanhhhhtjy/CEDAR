#!/usr/bin/env python3
"""Deterministic integrity and anonymity audit for the anonymous data archive."""

from __future__ import annotations

import argparse
import ast
import csv
import hashlib
import json
import re
import zipfile
from pathlib import Path
from typing import Any

TEXT_SUFFIXES = {".csv", ".json", ".jsonl", ".md", ".py", ".tsv", ".txt", ".yaml", ".yml"}
KNOWN_NAME_DIGESTS = {
    "8a38befe38c6332dc658bdb650f79ecfd20ce78d54f8721103274e874b31eed9",
    "737ca8c1cb898b370ec7f4a8e9d81ba81d62acefcda7732db27293665c1c6afc",
    "1614ceeec50a9336ebf690886caa747d6811c45d37086a3fa7b11c9e83926c6c",
}
IDENTITY_PATTERNS = {
    "private_unix_path": re.compile(r"/(?:data[14]|home|root|Users)/[A-Za-z0-9_.-]+"),
    "private_windows_path": re.compile(
        r"(?i)\b[A-Z]:[\\/](?:Users|ProgramData|Program Files|Windows|Temp|"
        r"[A-Za-z0-9_.-]{3,}[\\/])"
    ),
    "author_repository_url": re.compile(
        r"(?i)https?://[^\s\"']*(?:github\.com|raw\.githubusercontent\.com|"
        r"gitlab\.com|git\.overleaf\.com|overleaf\.com|anonymous\.4open\.science)"
    ),
    "credential": re.compile(
        r"(?i)\b(?:olp_[A-Za-z0-9]{20,}|ghp_[A-Za-z0-9]{20,}|"
        r"github_pat_[A-Za-z0-9_]{20,}|hf_[A-Za-z0-9]{20,}|sk-[A-Za-z0-9]{20,})\b"
    ),
    "compute_host": re.compile(r"(?i)\bautodl[-\w.]*\b"),
    "gpu_uuid": re.compile(r"(?i)\bGPU-[0-9a-f-]{20,}\b"),
    "private_email": re.compile(r"(?i)\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b"),
}


def contains_known_identity(text: str) -> bool:
    for token in re.findall(r"[A-Za-z0-9]+", text.lower()):
        if hashlib.sha256(token.encode("utf-8")).hexdigest() in KNOWN_NAME_DIGESTS:
            return True
    return False


def digest(path: Path) -> str:
    value = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            value.update(chunk)
    return value.hexdigest()


def audit_manifest(root: Path) -> list[dict[str, Any]]:
    manifest = root / "MANIFEST.csv"
    with manifest.open(encoding="utf-8", newline="") as handle:
        rows = list(csv.DictReader(handle))
    findings = []
    listed = set()
    for row in rows:
        relative = row["path"]
        listed.add(relative)
        path = root / relative
        if not path.is_file():
            findings.append({"kind": "manifest_missing_file", "path": relative})
            continue
        if int(row["bytes"]) != path.stat().st_size:
            findings.append({"kind": "manifest_size_mismatch", "path": relative})
        if row["sha256"] != digest(path):
            findings.append({"kind": "manifest_hash_mismatch", "path": relative})
    actual = {
        path.relative_to(root).as_posix()
        for path in root.rglob("*")
        if path.is_file() and path.name != "MANIFEST.csv"
    }
    for relative in sorted(actual - listed):
        findings.append({"kind": "manifest_unlisted_file", "path": relative})
    return findings


def audit_text(path: Path, relative: str) -> list[dict[str, Any]]:
    findings = []
    text = path.read_text(encoding="utf-8", errors="replace")
    for line_number, line in enumerate(text.splitlines(), 1):
        if contains_known_identity(line):
            findings.append(
                {"kind": "known_identity", "path": relative, "line": line_number}
            )
    for name, pattern in IDENTITY_PATTERNS.items():
        for match in pattern.finditer(text):
            line = text.count("\n", 0, match.start()) + 1
            findings.append({"kind": name, "path": relative, "line": line})
    suffix = path.suffix.lower()
    try:
        if suffix == ".json":
            json.loads(text)
        elif suffix == ".jsonl":
            for line_number, line in enumerate(text.splitlines(), 1):
                if line.strip():
                    json.loads(line)
        elif suffix == ".csv":
            with path.open(encoding="utf-8", newline="") as handle:
                list(csv.reader(handle))
        elif suffix == ".py":
            ast.parse(text, filename=relative)
    except Exception as exc:  # noqa: BLE001 - audit should retain every parser failure
        findings.append({"kind": "parse_failure", "path": relative, "detail": str(exc)})
    return findings


def audit_workbook(path: Path, relative: str) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=False)
        creator = workbook.properties.creator or ""
        modifier = workbook.properties.lastModifiedBy or ""
        findings = []
        if creator not in ("", "Anonymous Submission"):
            findings.append({"kind": "workbook_creator", "path": relative})
        if modifier not in ("", "Anonymous Submission"):
            findings.append({"kind": "workbook_modifier", "path": relative})
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if not isinstance(cell.value, str):
                        continue
                    if contains_known_identity(cell.value):
                        findings.append(
                            {
                                "kind": "workbook_known_identity",
                                "path": relative,
                                "sheet": sheet.title,
                                "cell": cell.coordinate,
                            }
                        )
                    for name, pattern in IDENTITY_PATTERNS.items():
                        if pattern.search(cell.value):
                            findings.append(
                                {
                                    "kind": f"workbook_{name}",
                                    "path": relative,
                                    "sheet": sheet.title,
                                    "cell": cell.coordinate,
                                }
                            )
        workbook.close()
        return findings
    except Exception as exc:  # noqa: BLE001
        return [{"kind": "workbook_parse_failure", "path": relative, "detail": str(exc)}]


def write_zip(root: Path, output: Path) -> None:
    if output.exists():
        output.unlink()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        for path in sorted(root.rglob("*")):
            if path.is_file():
                archive.write(path, path.relative_to(root.parent).as_posix())


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--package-root",
        type=Path,
        default=Path(__file__).resolve().parents[1],
    )
    parser.add_argument("--report", type=Path, default=None)
    parser.add_argument("--zip-output", type=Path, default=None)
    args = parser.parse_args()
    root = args.package_root.resolve()
    findings = audit_manifest(root)
    counts: dict[str, int] = {}
    for path in sorted(root.rglob("*")):
        if not path.is_file():
            continue
        suffix = path.suffix.lower()
        counts[suffix or "<none>"] = counts.get(suffix or "<none>", 0) + 1
        relative = path.relative_to(root).as_posix()
        if suffix in TEXT_SUFFIXES:
            findings.extend(audit_text(path, relative))
        elif suffix == ".xlsx":
            findings.extend(audit_workbook(path, relative))

    report = {
        "schema": "cedar.anonymous_supporting_package.audit.v2",
        "status": "passed" if not findings else "failed",
        "package_files": sum(counts.values()),
        "package_bytes": sum(path.stat().st_size for path in root.rglob("*") if path.is_file()),
        "suffix_counts": dict(sorted(counts.items())),
        "findings": findings,
    }
    report_path = args.report or root / "PACKAGE_AUDIT.json"
    report_path.write_text(
        json.dumps(report, indent=2, sort_keys=True) + "\n", encoding="utf-8", newline="\n"
    )
    # PACKAGE_AUDIT is intentionally created after checking MANIFEST; the caller
    # must rebuild MANIFEST once before the final manifest audit.
    if args.zip_output:
        write_zip(root, args.zip_output.resolve())
        report["zip_path"] = args.zip_output.name
        report["zip_bytes"] = args.zip_output.resolve().stat().st_size
    print(json.dumps(report, indent=2, sort_keys=True))
    return 0 if not findings else 1


if __name__ == "__main__":
    raise SystemExit(main())
