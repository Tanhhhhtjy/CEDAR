from __future__ import annotations

import copy
import inspect
import json
import shutil
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
HARDENING = ROOT / "datasets" / "m2_2b_formal" / "prover_generator_hardening_v1_1"
if str(HARDENING) not in sys.path:
    sys.path.insert(0, str(HARDENING))

from atomic_edit_candidate_enumerator import (  # noqa: E402
    enumerate_atomic_edit_candidates,
    evaluate_candidate_neighborhood,
)
from finalize_naturalization_review_v2 import (  # noqa: E402
    ReviewPacketError,
    finalize_review,
)
from hardening_common import (  # noqa: E402
    ERROR_CLASSES,
    MASTER_SEED,
    canonical_json_file_sha256,
    canonical_jsonl_sha256,
    canonical_sha256,
    file_sha256,
    stable_seed,
    write_json,
)
from naturalize_prover_generator_hardening_v1_1 import automatic_checks  # noqa: E402
from naturalize_prover_generator_hardening_v1_1_v2 import (  # noqa: E402
    WRAPPER_IDS,
    automatic_checks_v2,
    build_v2,
)
from proof_graph_builder import (  # noqa: E402
    FAMILY_BLUEPRINTS,
    assert_no_answer_channel,
    blueprint_by_id,
    proof_graph_builder,
)
from semantic_f_operators import derive_error_class, validate_f_certificate  # noqa: E402
from semantic_validator import semantic_validator  # noqa: E402

SENTINEL_FAMILIES = ["PGF-01", "PGF-02", "PGF-03", "PGF-04", "PGF-05", "PGF-06"]


def _case(family_id: str, error_class: str, variant: str = "easy"):
    blueprint = blueprint_by_id(family_id)
    seed = stable_seed(MASTER_SEED, "test", family_id, variant, error_class)
    built = proof_graph_builder(blueprint, variant, seed, blueprint["transition"], error_class)
    candidates = enumerate_atomic_edit_candidates(built)
    neighborhood = evaluate_candidate_neighborhood(built, candidates)
    return built, candidates, neighborhood


def test_builder_has_no_intended_edit_output_or_parameter() -> None:
    built, _, _ = _case("PGF-01", ERROR_CLASSES[0])
    assert_no_answer_channel(built)
    parameters = inspect.signature(proof_graph_builder).parameters
    assert "intended_edit" not in parameters
    assert "updated_world" not in parameters


def test_enumerator_has_no_answer_argument_and_is_deterministic() -> None:
    built, first, _ = _case("PGF-07", ERROR_CLASSES[0], "medium")
    second = enumerate_atomic_edit_candidates(built)
    assert inspect.signature(enumerate_atomic_edit_candidates).parameters.keys() == {
        "builder_output"
    }
    assert first == second
    assert 20 <= len(first) <= 100
    assert [item["rank"] for item in first] == list(range(1, len(first) + 1))


def test_all_48_neighborhoods_have_one_unique_edit_and_stable_order() -> None:
    counts = []
    class_variant_counts = {}
    for family_index, blueprint in enumerate(FAMILY_BLUEPRINTS):
        for variant_index, variant in enumerate(("easy", "medium")):
            for replicate_id in range(1, 4):
                error_class = ERROR_CLASSES[
                    (family_index * 3 + variant_index * 3 + replicate_id - 1) % 6
                ]
                class_variant_counts[(error_class, variant)] = (
                    class_variant_counts.get((error_class, variant), 0) + 1
                )
                seed = stable_seed(MASTER_SEED, blueprint["family_id"], variant, replicate_id)
                built = proof_graph_builder(
                    blueprint, variant, seed, blueprint["transition"], error_class
                )
                candidates = enumerate_atomic_edit_candidates(built)
                result = evaluate_candidate_neighborhood(built, candidates)
                assert result["success_count"] == 1
                assert result["successful_rank"] is not None
                assert candidates == enumerate_atomic_edit_candidates(built)
                counts.append(len(candidates))
    assert len(counts) == 48
    assert sum(value >= 20 for value in counts) / len(counts) >= 0.90
    assert set(class_variant_counts.values()) == {4}
    assert len(class_variant_counts) == 12


@pytest.mark.parametrize(
    ("family_id", "error_class"), zip(SENTINEL_FAMILIES, ERROR_CLASSES, strict=True)
)
def test_each_f_operator_positive_and_dual_gold(family_id: str, error_class: str) -> None:
    built, _, neighborhood = _case(family_id, error_class)
    result = semantic_validator(built, neighborhood)
    assert result["passed"]
    assert result["f_validation"]["derived_error_class"] == error_class
    assert result["initial_validation"]["agreement"]
    assert result["updated_validation"]["agreement"]


@pytest.mark.parametrize(
    ("family_id", "error_class"), zip(SENTINEL_FAMILIES, ERROR_CLASSES, strict=True)
)
def test_each_f_operator_rejects_unentailed_premise(family_id: str, error_class: str) -> None:
    built, _, _ = _case(family_id, error_class)
    context = built["fallacy_context"]
    forged = copy.deepcopy(context["premises"])
    forged[0]["predicate"] += "forged"
    with pytest.raises(ValueError, match="premises"):
        derive_error_class(built["initial_world"], forged, context["claimed_conclusion"])


def test_declared_error_class_cannot_forge_structural_class() -> None:
    built, _, neighborhood = _case("PGF-01", "cross_entity_transfer")
    result = semantic_validator(built, neighborhood)
    certificate = copy.deepcopy(result["f_certificate"])
    certificate["error_class"] = "unsupported_bridge"
    check = validate_f_certificate(
        built, certificate, result["selected_edit"], result["updated_world"]
    )
    assert not check["passed"]
    assert not check["checks"]["derived_class_matches_declared"]


def test_second_unsupported_downstream_step_fails() -> None:
    built, _, neighborhood = _case("PGF-02", "converse_error")
    result = semantic_validator(built, neighborhood)
    certificate = copy.deepcopy(result["f_certificate"])
    certificate["valid_downstream_steps"].append({"kind": "unsupported_extra_step"})
    check = validate_f_certificate(
        built, certificate, result["selected_edit"], result["updated_world"]
    )
    assert not check["passed"]
    assert not check["checks"]["no_second_unsupported_step"]


def test_easy_medium_formal_depth_contracts() -> None:
    for variant, low, high in (("easy", 2, 3), ("medium", 4, 6)):
        built, _, neighborhood = _case("PGF-08", "scope_transfer", variant)
        result = semantic_validator(built, neighborhood)
        assert result["passed"]
        assert low <= result["difficulty"]["actual_shortest_proof_depth"] <= high


def test_canonical_json_hash_ignores_key_order_and_line_endings(tmp_path: Path) -> None:
    first = tmp_path / "first.json"
    second = tmp_path / "second.json"
    first.write_bytes(b'{\r\n  "b": 2,\r\n  "a": 1\r\n}\r\n')
    write_json(second, {"a": 1, "b": 2})
    assert canonical_json_file_sha256(first) == canonical_json_file_sha256(second)


def test_no_prohibited_generic_fatal_step_in_v11_source() -> None:
    prohibited = "unsupported" + "_atomic_state_claim"
    for path in HARDENING.glob("*.py"):
        assert prohibited not in path.read_text(encoding="utf-8")


def test_naturalization12_passes_surface_and_reverse_mapping_checks() -> None:
    rows = [
        json.loads(line)
        for line in (HARDENING / "naturalization12.jsonl").read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]
    assert len(rows) == 12
    assert automatic_checks(rows)["passed"]
    cells = {(row["internal_error_class"], row["variant"]) for row in rows}
    assert len(cells) == 12


def test_review_workbook_hides_internals_and_leaves_human_fields_blank() -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(HARDENING / "naturalization_review12_bilingual.xlsx", data_only=False)
    assert workbook.sheetnames == ["自然化审核", "审核说明"]
    sheet = workbook["自然化审核"]
    surface = "\n".join(
        str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value is not None
    )
    for error_class in ERROR_CLASSES:
        assert error_class not in surface
    assert "source_attempt_id" not in surface
    for row in range(6, 18):
        assert all(sheet.cell(row=row, column=column).value is None for column in range(4, 9))
        assert sheet.cell(row=row, column=13).value is None


def test_manifest_has_no_absolute_paths_and_binds_all_48_attempts() -> None:
    manifest = json.loads((HARDENING / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["gpu_qwen_capping_run"] is False
    assert len(manifest["formal_attempts"]) == 48
    assert not manifest["failures"]
    assert manifest["naturalization_amendment"] == "02"
    assert manifest["protected_baseline_commit"].startswith("06fd188")
    assert manifest["naturalization12_v2_canonical_sha256"]
    assert manifest["workbook_v2_canonical_content_sha256"]
    assert manifest["human_review_finalizer_sha256"]
    assert manifest["completed_workbook_raw_sha256"]
    assert manifest["completed_workbook_review_content_canonical_sha256"]
    assert manifest["human_review_result_canonical_sha256"]
    assert all(":" not in entry["path"] for entry in manifest["files"])
    assert all(not entry["path"].endswith("/_dump.json") for entry in manifest["files"])
    assert "C:\\" not in json.dumps(manifest)


def test_amendment02_manifest_hashes_match_v2_delivery() -> None:
    manifest = json.loads((HARDENING / "manifest.json").read_text(encoding="utf-8"))
    entries = {Path(entry["path"]).name: entry for entry in manifest["files"]}
    required = {
        "naturalization12_v2.jsonl",
        "naturalization_automatic_checks_v2.json",
        "naturalization_determinism_v2.json",
        "renderer_revision_diff.json",
        "naturalization_review12_bilingual_v2.xlsx",
        "finalize_naturalization_review_v2.py",
    }
    assert required.issubset(entries)
    for filename in required:
        path = HARDENING / filename
        assert entries[filename]["raw_file_sha256"] == file_sha256(path)
        if path.suffix == ".jsonl":
            expected = canonical_jsonl_sha256(path)
        elif path.suffix == ".json":
            expected = canonical_json_file_sha256(path)
        elif path.suffix == ".xlsx":
            expected = manifest["workbook_v2_canonical_content_sha256"]
        else:
            normalized = path.read_text(encoding="utf-8").replace("\r\n", "\n")
            expected = canonical_sha256({"normalized_text": normalized})
        assert entries[filename]["canonical_content_sha256"] == expected


def test_final_gate_stops_at_real_human_review() -> None:
    gate = json.loads(
        (ROOT / "explore_outputs/prover_generator_hardening_v1_1/FINAL_GATE.json").read_text(
            encoding="utf-8"
        )
    )
    assert gate["gate"] == "C"
    assert gate["symbolic_status"] == "symbolic_generator_passed"
    assert gate["human_review_status"] == "complete"
    assert gate["naturalization_status"] == "naturalization_failed"
    assert gate["failed_review_ids"] == [f"NAT-{index:02d}" for index in range(1, 13)]
    assert gate["step_1_complete"] is False
    assert gate["may_prepare_development_bank_protocol"] is False


def test_analyzer_preserves_finalized_gate_c() -> None:
    result = json.loads((HARDENING / "human_review_result.json").read_text(encoding="utf-8"))
    assert result["status"] == "failed"
    assert result["passed_row_count"] == 0
    assert all(row["language_naturalness"] == 3 for row in result["rows"])
    assert all(row["condition_identity_leakage"] == "yes" for row in result["rows"])


def _v2_rows() -> list[dict]:
    return [
        json.loads(line)
        for line in (HARDENING / "naturalization12_v2.jsonl")
        .read_text(encoding="utf-8")
        .splitlines()
        if line.strip()
    ]


def _synthetic_completed_workbook(
    tmp_path: Path, *, failed_row: int | None = None, tamper: bool = False
) -> Path:
    """Create an uncommitted test fixture; this is never treated as human review evidence."""

    from openpyxl import load_workbook

    source = HARDENING / "naturalization_review12_bilingual_v2.xlsx"
    destination = tmp_path / "synthetic_completed_review.xlsx"
    shutil.copy2(source, destination)
    workbook = load_workbook(destination)
    sheet = workbook["自然化审核v2"]
    for row in range(6, 18):
        sheet.cell(row=row, column=4, value=5)
        sheet.cell(row=row, column=5, value="yes")
        sheet.cell(row=row, column=6, value="yes")
        sheet.cell(row=row, column=7, value="no")
        sheet.cell(row=row, column=8, value="no")
        sheet.cell(row=row, column=9, value="yes")
    if failed_row is not None:
        sheet.cell(row=failed_row, column=8, value="yes")
    if tamper:
        sheet.cell(row=6, column=2, value=str(sheet.cell(row=6, column=2).value) + " altered")
    workbook.save(destination)
    return destination


def test_amendment02_renderer_contract_and_leakage_report() -> None:
    rows = _v2_rows()
    report = automatic_checks_v2(rows)
    assert len(rows) == 12
    assert report["passed"]
    assert all("treats the claim that" in row["english_f_text"] for row in rows)
    assert all(
        row["reverse_mapping"]["selected_edit"]["operation"] != "replace"
        or (
            "replace the record ‘" in row["english_v_text"]
            and "’ with ‘" in row["english_v_text"]
            and "replace the record ‘" in row["english_f_text"]
        )
        for row in rows
    )
    leakage = report["condition_leakage"]
    assert leakage["high_risk_exclusive_phrase_count"] == 0
    assert set(leakage["template_distribution"]["overall"]) == set(WRAPPER_IDS)
    assert set(leakage["template_distribution"]) == {
        "overall",
        "by_family",
        "by_difficulty",
        "by_f_class",
    }


def test_amendment02_hash_only_template_assignment() -> None:
    rows = _v2_rows()
    assert all(row["template_selector_input"] == "source_symbolic_sha256_only" for row in rows)
    assert all(
        field not in row["template_selector_input"]
        for row in rows
        for field in ("family", "variant", "error", "gold", "transition")
    )


def test_amendment02_protected_symbolic_artifacts_unchanged() -> None:
    diff = json.loads((HARDENING / "renderer_revision_diff.json").read_text(encoding="utf-8"))
    assert diff["all_source_hashes_unchanged"]
    assert diff["all_base_prompts_unchanged"]
    assert diff["protected_symbolic_artifacts"]["all_unchanged"]
    assert diff["protected_symbolic_artifacts"]["protected_file_count"] > 300


def test_amendment02_blank_workbook_has_new_field_and_source_hashes() -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(
        HARDENING / "naturalization_review12_bilingual_v2.xlsx", data_only=False
    )
    sheet = workbook["自然化审核v2"]
    assert sheet.cell(row=5, column=8).value == "条件身份泄漏"
    assert sheet.column_dimensions["O"].hidden
    source_hashes = {row["source_symbolic_sha256"] for row in _v2_rows()}
    assert {sheet.cell(row=row, column=15).value for row in range(6, 18)} == source_hashes
    for row in range(6, 18):
        assert all(sheet.cell(row=row, column=column).value is None for column in range(4, 10))
        assert sheet.cell(row=row, column=14).value is None


def test_amendment02_two_canonical_builds_match(tmp_path: Path) -> None:
    first = tmp_path / "first"
    second = tmp_path / "second"
    first.mkdir()
    second.mkdir()
    report_first = build_v2(HARDENING, first)
    report_second = build_v2(HARDENING, second)
    assert (
        report_first["naturalization_canonical_sha256"]
        == report_second["naturalization_canonical_sha256"]
    )
    assert (
        report_first["workbook_canonical_content_sha256"]
        == report_second["workbook_canonical_content_sha256"]
    )
    for filename in ("naturalization12_v2.jsonl", "renderer_revision_diff.json"):
        assert (first / filename).read_bytes() == (second / filename).read_bytes()
    first_checks = json.loads(
        (first / "naturalization_automatic_checks_v2.json").read_text(encoding="utf-8")
    )
    second_checks = json.loads(
        (second / "naturalization_automatic_checks_v2.json").read_text(encoding="utf-8")
    )
    first_checks.pop("workbook_raw_file_sha256")
    second_checks.pop("workbook_raw_file_sha256")
    assert first_checks == second_checks


def test_finalizer_rejects_incomplete_blank_packet(tmp_path: Path) -> None:
    with pytest.raises(ReviewPacketError, match="must be an integer"):
        finalize_review(
            HARDENING / "naturalization_review12_bilingual_v2.xlsx",
            result_path=tmp_path / "result.json",
            final_gate_path=tmp_path / "gate.json",
        )
    assert not (tmp_path / "result.json").exists()
    assert not (tmp_path / "gate.json").exists()


def test_finalizer_rejects_immutable_text_change(tmp_path: Path) -> None:
    completed = _synthetic_completed_workbook(tmp_path, tamper=True)
    with pytest.raises(ReviewPacketError, match="immutable cell changed"):
        finalize_review(
            completed,
            result_path=tmp_path / "result.json",
            final_gate_path=tmp_path / "gate.json",
        )


def test_finalizer_all_rows_pass_produces_gate_a(tmp_path: Path) -> None:
    completed = _synthetic_completed_workbook(tmp_path)
    output = finalize_review(
        completed,
        result_path=tmp_path / "result.json",
        final_gate_path=tmp_path / "gate.json",
    )
    assert output["result"]["passed_row_count"] == 12
    assert output["result"]["agent_filled_human_ratings"] is False
    assert output["final_gate"]["gate"] == "A"
    assert output["final_gate"]["step_1_complete"] is True


def test_finalizer_one_item_failure_produces_gate_c(tmp_path: Path) -> None:
    completed = _synthetic_completed_workbook(tmp_path, failed_row=9)
    output = finalize_review(
        completed,
        result_path=tmp_path / "result.json",
        final_gate_path=tmp_path / "gate.json",
    )
    assert output["result"]["passed_row_count"] == 11
    assert output["result"]["failed_review_ids"] == ["NAT-04"]
    assert output["final_gate"]["gate"] == "C"
    assert output["final_gate"]["renderer_revision_only"] is True
