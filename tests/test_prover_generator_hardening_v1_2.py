from __future__ import annotations

import copy
import json
import shutil
import sys
from collections import Counter
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
HARDENING = ROOT / "datasets/m2_2b_formal/prover_generator_hardening_v1_2"
if str(HARDENING) not in sys.path:
    sys.path.insert(0, str(HARDENING))

from build_prover_generator_hardening_v1_2 import validate_row  # noqa: E402
from finalize_paired_justification_review import (  # noqa: E402
    ReviewPacketError,
    finalize,
)
from paired_justification_common import (  # noqa: E402
    ERROR_CLASSES,
    _render,
    canonical_json_file_sha256,
    canonical_jsonl_sha256,
    file_sha256,
    read_jsonl,
    select_source_rows,
    stable_text_sha256,
)
from validate_prover_generator_hardening_v1_2 import validate_formal  # noqa: E402


def _rows() -> list[dict]:
    return read_jsonl(HARDENING / "paired_justification_sentinels12.jsonl")


def test_exact_sentinel_matrix() -> None:
    rows = _rows()
    assert len(rows) == 12
    assert Counter((row["error_class"], row["variant"]) for row in rows) == Counter(
        {(name, variant): 1 for name in ERROR_CLASSES for variant in ("easy", "medium")}
    )
    assert all(row["sentinel"] is True for row in rows)


def test_pair_identity_fields_and_edit_are_exactly_equal() -> None:
    for row in _rows():
        v = row["pair_contract"]["v"]
        f = row["pair_contract"]["f"]
        assert v == f
        assert v["slot_skeleton_sha256"] == row["surface"]["slot_skeleton_sha256"]
        assert v["proposed_edit_text_en"] == row["surface"]["proposed_edit_text_en"]


def test_all_stored_and_recomputed_automatic_checks_pass() -> None:
    for row in _rows():
        assert row["automatic_validation"]["passed"]
        assert validate_row(row)["passed"]


def test_independent_formal_validation_passes_every_pair() -> None:
    for row in _rows():
        result = validate_formal(row)
        assert result["passed"], (row["pair_id"], result)


def test_error_class_cannot_be_forged_by_string() -> None:
    row = copy.deepcopy(_rows()[0])
    row["formal"]["f"]["certificate"]["declared_error_class"] = "converse_error"
    assert not validate_formal(row)["passed"]


def test_second_unsupported_step_is_rejected() -> None:
    row = copy.deepcopy(_rows()[0])
    bogus = copy.deepcopy(row["formal"]["query"])
    bogus["predicate"] = "unsupported_second_step"
    row["formal"]["f"]["certificate"]["downstream_literals"].append(bogus)
    assert not validate_formal(row)["passed"]


def test_proof_depth_gradient_and_unique_v_proof() -> None:
    rows = _rows()
    for name in ERROR_CLASSES:
        easy = next(row for row in rows if row["error_class"] == name and row["variant"] == "easy")
        medium = next(
            row for row in rows if row["error_class"] == name and row["variant"] == "medium"
        )
        easy_depth = easy["automatic_validation"]["v_warrant_proof_depth"]
        medium_depth = medium["automatic_validation"]["v_warrant_proof_depth"]
        assert 2 <= easy_depth <= 3
        assert 4 <= medium_depth <= 6
        assert medium_depth > easy_depth
        assert easy["formal"]["v"]["proof_count"] == 1
        assert medium["formal"]["v"]["proof_count"] == 1


def test_surface_language_contract() -> None:
    for row in _rows():
        result = row["automatic_validation"]
        assert 0.90 <= result["v_f_word_count_ratio"] <= 1.10
        assert result["checks"]["coreference"]
        assert result["checks"]["forbidden_language_absent"]
        assert result["checks"]["no_encoding_artifacts"]
        assert "that same item" in (
            row["surface"]["english_v_text"] + row["surface"]["english_f_text"]
        ) or "an item" not in (
            row["surface"]["english_v_text"] + row["surface"]["english_f_text"]
        )


def test_base_prompt_coreference_repair_is_limited_to_pj11_and_pj12() -> None:
    rows = _rows()
    sources = select_source_rows()
    for index, (row, source) in enumerate(zip(rows, sources, strict=True), start=1):
        frozen = _render(source, index)
        english = row["surface"]["english_base_prompt"]
        chinese = row["surface"]["chinese_base_prompt"]
        if index in {11, 12}:
            assert english != frozen["english_base_prompt"]
            assert chinese != frozen["chinese_base_prompt"]
            assert "then an item has status" not in english
            assert "then that same item has status" in english
            assert "那么某个对象具有" not in chinese
            assert "那么该同一对象具有" in chinese
        else:
            assert english == frozen["english_base_prompt"]
            assert chinese == frozen["chinese_base_prompt"]


def test_automatic_and_determinism_reports_pass() -> None:
    automatic = json.loads(
        (HARDENING / "automatic_validation_report.json").read_text(encoding="utf-8")
    )
    deterministic = json.loads(
        (HARDENING / "deterministic_rebuild_report.json").read_text(encoding="utf-8")
    )
    assert automatic["passed"]
    assert deterministic["passed"]
    assert all(deterministic["compared"].values())


def test_blank_workbook_hides_internals_and_ratings() -> None:
    workbook = load_workbook(
        HARDENING / "paired_justification_review12_bilingual.xlsx", data_only=False
    )
    assert workbook.sheetnames == ["自然化审核v1.2", "审核说明"]
    sheet = workbook["自然化审核v1.2"]
    assert all(
        sheet.cell(row=row, column=column).value is None
        for row in range(6, 18)
        for column in range(8, 15)
    )
    visible = "\n".join(
        str(sheet.cell(row=row, column=column).value or "")
        for row in range(1, 18)
        for column in range(1, 15)
    ).lower()
    assert all(name not in visible for name in ERROR_CLASSES)
    assert "revision_warranted" not in visible
    assert sheet.column_dimensions["O"].hidden


def test_manifest_hashes_and_boundaries() -> None:
    manifest = json.loads((HARDENING / "manifest.json").read_text(encoding="utf-8"))
    entries = {Path(item["path"]).name: item for item in manifest["files"]}
    assert manifest["sentinel_count"] == 12
    assert manifest["gpu_qwen_capping_run"] is False
    assert manifest["development_bank_built"] is False
    assert manifest["human_review_status"] == "awaiting_human_naturalization_review"
    assert "C:\\" not in json.dumps(manifest)
    assert entries["paired_justification_sentinels12.jsonl"][
        "canonical_content_sha256"
    ] == canonical_jsonl_sha256(HARDENING / "paired_justification_sentinels12.jsonl")
    assert entries["automatic_validation_report.json"][
        "canonical_content_sha256"
    ] == canonical_json_file_sha256(HARDENING / "automatic_validation_report.json")
    for entry in manifest["files"]:
        path = ROOT / entry["path"]
        assert entry["raw_file_sha256"] == file_sha256(path)
        if path.suffix == ".jsonl":
            expected = canonical_jsonl_sha256(path)
        elif path.suffix == ".json":
            expected = canonical_json_file_sha256(path)
        elif path.suffix == ".xlsx":
            expected = manifest["workbook_canonical_content_sha256"]
        else:
            expected = stable_text_sha256(path.read_text(encoding="utf-8"))
        assert entry["canonical_content_sha256"] == expected


def test_final_gate_stops_for_real_human_review() -> None:
    gate = json.loads(
        (ROOT / "explore_outputs/prover_generator_hardening_v1_2/FINAL_GATE.json").read_text(
            encoding="utf-8"
        )
    )
    assert gate["gate"] == "B"
    assert gate["status"] == "awaiting_human_naturalization_review"
    assert gate["step_1_complete"] is False
    assert gate["may_prepare_development_bank_protocol"] is False


def test_v1_1_gate_c_remains_frozen() -> None:
    gate = json.loads(
        (ROOT / "explore_outputs/prover_generator_hardening_v1_1/FINAL_GATE.json").read_text(
            encoding="utf-8"
        )
    )
    assert gate["gate"] == "C"
    assert gate["symbolic_status"] == "symbolic_generator_passed"


def _completed_workbook(tmp_path: Path, *, fail_row: int | None = None) -> Path:
    source = HARDENING / "paired_justification_review12_bilingual.xlsx"
    target = tmp_path / "completed.xlsx"
    shutil.copy2(source, target)
    workbook = load_workbook(target)
    sheet = workbook["自然化审核v1.2"]
    for row in range(6, 18):
        sheet.cell(row, 8, 5)
        sheet.cell(row, 9, "yes")
        sheet.cell(row, 10, "yes")
        sheet.cell(row, 11, "no")
        sheet.cell(row, 12, "no")
        sheet.cell(row, 13, "yes")
    if fail_row is not None:
        sheet.cell(fail_row, 12, "yes")
    workbook.save(target)
    return target


def test_finalizer_derives_gate_a_and_c_without_editing_workbook(tmp_path: Path) -> None:
    passed = _completed_workbook(tmp_path)
    passed_raw = passed.read_bytes()
    output = finalize(
        passed,
        result_path=tmp_path / "passed_result.json",
        gate_path=tmp_path / "passed_gate.json",
    )
    assert output["gate"]["gate"] == "A"
    assert output["result"]["passed_row_count"] == 12
    assert passed.read_bytes() == passed_raw

    failed = _completed_workbook(tmp_path, fail_row=6)
    output = finalize(
        failed,
        result_path=tmp_path / "failed_result.json",
        gate_path=tmp_path / "failed_gate.json",
    )
    assert output["gate"]["gate"] == "C"
    assert output["result"]["failed_review_ids"] == ["PJ-01"]


def test_finalizer_rejects_immutable_text_change(tmp_path: Path) -> None:
    completed = _completed_workbook(tmp_path)
    workbook = load_workbook(completed)
    workbook["自然化审核v1.2"]["D6"] = "tampered"
    workbook.save(completed)
    with pytest.raises(ReviewPacketError, match="immutable cell changed"):
        finalize(
            completed,
            result_path=tmp_path / "result.json",
            gate_path=tmp_path / "gate.json",
        )
